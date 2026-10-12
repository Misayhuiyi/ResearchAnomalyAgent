"""Reproducibly assemble a new candidate, never overwrite an earlier release.

This packages executable source and truthful operator documentation. It does not
certify platform integration, model quality, Docker execution or human acceptance.
Run with the bundled document-authoring Python runtime.
"""
import csv
import ast
import hashlib
import html
import json
import shutil
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'mvp-demo'
NAME='PaperAudit'
TITLE='论文科研数据异常筛查应用说明'
INPUTS=[
 ['Materials','文件夹上传','作用为提供待检查材料。必填，同一论文的 1 至 20 份 CSV、TSV、Excel、文字 PDF 或 Word 文件，单份不超过 20 MB，合计不超过 100 MB。文件命名无需固定前缀，匹配使用完整原名，重名按不区分大小写校验；不根据文件名猜测论文归属，由用户确认后批量处理。注意仅平级文件，不允许链接、嵌套目录、压缩包及重名文件。'],
 ['Profile.txt','单项选择','作用为选择检查范围。可选，默认 review 标准检查；triage 为快速，forensic 为全面。格式为 UTF-8 单行文本，填写一个选项。注意检查范围仍以实际报告为准。'],
 ['Checks.txt','多项选择','作用为选择额外核对项目。可选，默认空。UTF-8 文本，每行一个项目，以换行分隔。由工作台中文选项生成；可选项见使用说明中的核对项目表。注意只有满足计算前提才可勾选。'],
 ['StudyNote.txt','单行文本','作用为说明分组和单位。可选，默认空，格式为 UTF-8 文本，最多 2000 字。涉及文件、分组或列名时，名称匹配须与原材料完全一致。注意它仅为用户说明，不修改原始证据。'],
]
OUTPUTS=[
 ['01.Report/','用于阅读和共享同一冻结版本的 Word、PDF、HTML 报告。包含问题、位置、原值、覆盖范围与待复核数量；数值线索不等于科研不端。批处理尚未人工复核，也未进行付费内容核对。'],
 ['02.Details/','用于逐项核对的 Excel 明细。字段包括问题、文件、行列位置、原始值和复核状态；待复核表示尚未由人作出结论。'],
 ['03.Evidence/','用于回查来源的 evidence.json。material_sources 为材料来源，findings 为数值证据，study_note 为用户说明；不将用户说明视为已证实事实。须按材料所属用户控制访问。'],
 ['Process/','用于恢复和复现的结构化 result.json。findings 为线索，scan_status 为读取范围；partial 或 failed 不能解释成无异常，不作为普通用户主入口。'],
 ['PipeLine/','用于内部复现的步骤快照。目录字段 01_Scan、02_Report 分别表示扫描与报告组装，各自 TempData 为实际输入，Result 为输出；仅本批文件参与归档。'],
 ['FinalResults/','本次检查概要。summary.json 的 scan_status 表示读取范围，finding_count 是待核对项数量，reviewed 与 model_used 说明人工复核和模型是否执行。'],
 ['Reports/','用于打开报告的 index.html。链接字段分别指向同一批次的 Word、PDF、HTML 和 Excel，不能混用不同批次的结论。'],
 ['Figures/','用于说明图件范围。本应用不生成统计图，无统计图例或坐标判读；数值以证据表为准。'],
 ['PipelineStatus.csv','用于核查各步骤状态。step 为步骤，status 为完成情况，message 为范围说明；读取失败不等于未发现问题。'],
 ['InputOutputManifest.csv','用于核对导出文件。字段含义：path 为相对路径，sha256 为内容校验，bytes 为文件大小；校验不一致时不得将文件视为同一版本。'],
 ['Running_log.txt','用于后台诊断的文本记录。判读内容为处理是否结束、模型与人工复核是否执行；不向普通用户展示堆栈、密钥或模型原话。'],
 ['session_info.txt','用于管理员核查环境。字段为 Python 版本和操作系统，不能据此证明容器或平台已通过运行验收。'],
]
SECTIONS=[
 ('应用信息',['应用名称：论文科研数据异常筛查。应用 ID：PaperAudit。应用开发者：吴梓豪。',
   '本应用用于将论文材料中的数值异常整理为可逐项核对的线索。先检查原始数值和位置，再由用户结合原文复核。结果不能单独认定科研不端。']),
 ('应用功能',['支持同一篇论文多份材料的数值检查，显示原始值及文件位置。支持有明确列名和计算前提的统计关系核对。读取不完整、条件不足和未执行的项目单独说明。工作台可保存人工复核，防止旧页面覆盖新意见。补充材料形成新版本，旧结果和旧报告保留。用户可为单条已发现线索请求 DeepSeek V4.1 Flash 中文辅助解释；每条先确认费用，只发送少量证据摘要，模型建议不代替人工结论。']),
 ('画图说明',['本应用不生成统计推断图。工作台以原始数据表和证据对照显示结果，不需要图片尺寸或颜色参数。']),
 ('应用场景',['研究人员自查论文数据，编辑或科研管理人员核对作者提供的材料。上传前应确认有权处理材料，并获得向外部模型发送相关片段的必要授权。扫描 PDF 和图像表格不在本版范围。']),
 ('使用说明',['在工作台选择材料或查找公开下载链接。确认材料属于同一篇论文，并仅勾选满足前提的核对项目。检查结束后先查看读取范围，再逐项核对原始数据。',
   '选择一条线索，核对原始数值和数学关系。需要辅助说明时，查看本条将发送的摘要及费用上限，再确认生成解释。解释包含可能的正常原因和人工核验步骤；调用未完成时仍可查看原始检查结果。',
   '填写人工复核后生成报告。四种格式来自同一冻结版本。补充或更正材料时建立新版本，不继承旧版本的复核结论。',
   '批处理入口只执行数值检查和初始报告；交互式费用确认、人工复核和历史记录通过工作台完成。']),
 ('参考文献',['PaperConan 0.8.5 软件包：本项目调用的数值异常筛查引擎。概要为在输入数据中查找可进一步核对的数值关系；本项目补充原始证据定位、覆盖说明和人工复核。具体算法和适用条件以随包规则目录及本项目 Spec 为准。',
   '《人人都能成为“耿同学”：开源工具 PaperConan 介绍》。文章摘要：通过确定性数值检查定位源数据线索，再结合原表、图注和研究方法核对共享对照、单位换算等合理解释。与本应用的关系是将计算、原文依据和人工复核分开，不据统计信号直接判断作者意图。原文：https://mp.weixin.qq.com/s/G6uB_0PXjywlrKOI93PvBA 。',
   'DeepSeek API 文档 Chat Completion：https://api-docs.deepseek.com/api/create-chat-completion/ 。文档描述请求、响应和用量字段，本项目用于受费用授权约束的论文内容核对适配。接口可用性不等同于模型效果已获验收。']),
]


def put(path,text):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text(text,encoding='utf-8')


def application_module(text):
    """Package app as an imported service, preserving its actual function bodies.

    Only the development CLI is omitted; the formal no-argument entry remains
    PaperAudit.py and the interactive entry is harness/Workbench.py.
    """
    lines=text.splitlines(keepends=True)
    remove=set()
    for node in ast.parse(text).body:
        if ((isinstance(node,ast.Import) and any(alias.name=='argparse' for alias in node.names))
                or (isinstance(node,ast.If) and ast.unparse(node.test)=="__name__ == '__main__'")):
            remove.update(range(node.lineno-1,node.end_lineno))
    return ''.join(line for index,line in enumerate(lines) if index not in remove)


def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=value
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=value;cells[i].width=Inches(widths[i])
    for row in t.rows:
        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                border=OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:border.set(qn('w:'+key),value)
                borders.append(border)
            tcpr.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(5);paragraph.paragraph_format.space_before=Pt(5)
    for cell in t.rows[0].cells:
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F6');cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph()


def documents(root):
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=section.right_margin=Inches(.75)
    section.top_margin=section.bottom_margin=Inches(.65)
    for name,size in [('Normal',11),('Title',22),('Heading 1',15)]:
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        ppr=style.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
    doc.core_properties.author='吴梓豪';doc.core_properties.title=TITLE
    doc.add_paragraph(TITLE,'Title')
    parts=['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>'+TITLE+'</title><style>body{font:16px/1.8 Microsoft YaHei,sans-serif;color:#182c3e;max-width:1000px;margin:40px auto;padding:20px}h1,h2{color:#000}table{border-collapse:collapse;width:100%}td,th{border:1px solid #cbd5e1;padding:12px;text-align:left;vertical-align:top}</style><h1>'+TITLE+'</h1>']
    markdown=['# '+TITLE,'']
    for title,paragraphs in SECTIONS:
        if title=='画图说明':
            for label,headers,rows,widths in [('应用输入',['应用输入（参数）','参数类型','参数说明'],INPUTS,[1.3,1.0,4.7]),('应用输出',['结果文件','结果说明'],OUTPUTS,[1.8,5.2])]:
                doc.add_heading(label,1);table(doc,headers,rows,widths)
                parts.append('<h2>'+label+'</h2><table><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in headers)+'</tr>'+''.join('<tr>'+''.join('<td>'+html.escape(c)+'</td>' for c in row)+'</tr>' for row in rows)+'</table>')
                markdown.extend(['## '+label,'','| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |',*['| '+' | '.join(row)+' |' for row in rows],''])
        doc.add_heading(title,1);parts.append('<h2>'+title+'</h2>');markdown.extend(['## '+title,''])
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph);parts.append('<p>'+html.escape(paragraph)+'</p>');markdown.extend([paragraph,''])
    doc.save(root/(NAME+' 应用说明.docx'))
    put(root/(NAME+' 应用说明.html'),''.join(parts)+'</html>')
    put(root/(NAME+' 应用说明.md'),'\n'.join(markdown))
    # Code-drawn UI-style cover, no reused photograph or fabricated result.
    cover=Image.new('RGB',(500,320),'#f4f8fd');draw=ImageDraw.Draw(cover)
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',26);small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',16)
    draw.text((30,28),'论文科研数据异常筛查',font=font,fill='#142e4a')
    draw.text((30,74),'原始数值   ·   原文依据   ·   人工复核',font=small,fill='#526781')
    draw.rounded_rectangle((30,116,470,278),radius=12,fill='white',outline='#cbdbee',width=2)
    for y,text in [(136,'选择同一篇论文的材料'),(178,'检查数值并查看原始位置'),(220,'保存复核并生成报告')]:
        draw.ellipse((49,y+5,61,y+17),fill='#367fd4');draw.text((79,y),text,font=small,fill='#18344b')
    cover.save(root/(NAME+' 应用封面.png'))


def main():
    stamp=datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')
    target=ROOT/'delivery'/(NAME+'-'+stamp);target.mkdir(parents=True,exist_ok=False)
    coding=target/(NAME+'_Coding')
    shutil.copytree(ROOT/'delivery-src'/'Code',coding/'Code')
    core=coding/'Code'/'DataCode'
    for source in SOURCE.glob('*.py'):
        if source.name=='app.py':put(core/source.name,application_module(source.read_text(encoding='utf-8')))
        else:shutil.copy2(source,core/source.name)
    for folder in ['static','samples']:shutil.copytree(SOURCE/folder,core/folder)
    for name in ['Data','TempData','Result']:(coding/name).mkdir(exist_ok=True)
    inputs=coding/'TempData';(inputs/'Materials').mkdir()
    shutil.copy2(SOURCE/'samples'/'fixed_shift_demo.csv',inputs/'Materials'/'fixed_shift_demo.csv')
    for name,value in [('Profile.txt','review'),('Checks.txt',''),('StudyNote.txt','')]:put(inputs/name,value)
    shutil.copytree(inputs,coding/'Data'/'Demo'/'TempData')
    shutil.copy2(SOURCE/'requirements.txt',coding/'Data'/'requirements.txt')
    shutil.copy2(ROOT/'THIRD_PARTY_NOTICES.md',coding/'Data'/'THIRD_PARTY_NOTICES.md')
    put(coding/'Data'/'模型配置说明.md','# 模型配置说明\n\n优先读取 DEEPSEEK_API_KEY、DEEPSEEK_BASE_URL、DEEPSEEK_MODEL 环境变量。官方地址为 https://api.deepseek.com，模型标识 deepseek-flash 对应 DeepSeek V4.1 Flash。Windows 管理员可在交互终端运行 Coding/Code/DataCode/OperatorConfigureModel.py，隐藏输入并加密保存到当前用户的本机目录；不要将密钥放入交付包。每条线索在工作台确认保守费用上限后，才发送最多 3000 字节的证据摘要。批处理不会调用模型。详见 https://api-docs.deepseek.com/quick_start/pricing/ 。\n')
    put(coding/'Data'/'本地运维说明.md','# 本地运维说明\n\n工作台仅供本机单用户使用，状态位于 Agent/Result/WorkbenchState。同一状态目录只能运行一个实例。停止服务后应备份完整状态目录，包含输入材料证据、扫描、复核和费用记录。模型超时不自动重发；服务方账单经人工核对后，可由管理员通过 Coding/Code/DataCode/OperatorReconcile.py 对账。平台身份、存储、论坛和容器调度由平台方接入。\n')
    shutil.copy2(ROOT/'tools'/'reconcile_model_bill.py',core/'OperatorReconcile.py')
    shutil.copy2(ROOT/'tools'/'configure_local_model.py',core/'OperatorConfigureModel.py')
    shutil.copy2(SOURCE/'model-config.example.json',coding/'Data'/'model-config.example.json')
    def constant(filename,name):
        tree=ast.parse((SOURCE/filename).read_text(encoding='utf-8'))
        return next(ast.literal_eval(node.value) for node in tree.body if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id==name for t in node.targets))
    rules=constant('statistical_checks.py','RULES')
    rules.update({'patterns_'+k:v for k,v in constant('pattern_checks.py','MODES').items()})
    rules_text='# 使用说明中的核对项目表\n\n工作台显示中文名称。以下项目名仅供管理员生成 Checks.txt；每行一项，不得在未确认计算前提时批量开启。\n\n| 项目名 | 中文名称 | 前提 |\n| --- | --- | --- |\n'+'\n'.join('| '+' | '.join([key,value['name'],value['assumption']])+' |' for key,value in rules.items())+'\n'
    put(target/'使用说明中的核对项目表.md',rules_text)
    pipeline=coding/'Data'/'PipeLine'
    contract={'workflow_id':NAME,'entry':'PaperAudit.py','callers':['PaperConan.py'],
              'platform_inputs':['Materials','Profile.txt','Checks.txt','StudyNote.txt'],
              'steps':['01_Scan','02_Report'],'final_containers':['PipeLine','FinalResults','Figures','Reports','01.Report','02.Details','03.Evidence','Process']}
    put(pipeline/'WorkflowContract.json',json.dumps(contract,ensure_ascii=False,indent=2))
    put(pipeline/'PipelineConfig.json',json.dumps({'automatic_retries':0,'paid_model_calls':False},indent=2))
    put(pipeline/'StepIOMap.csv','source_step,source_file,target_step,target_file,required,description\n01_Scan,result.json,02_Report,result.json,true,冻结数值检查结果后导出报告\n')
    put(pipeline/'AppManifest.csv','step_id,application_id,caller,image_or_version,description\n01_Scan,PaperConan,PaperConan.py,0.8.5,数值检查\n02_Report,PaperAudit,,1,本应用内置同版本报告组装\n')
    for step in contract['steps']:
        for folder in ['TempData','Result']:(pipeline/step/folder).mkdir(parents=True)
    put(coding/'Data'/'Dockerfile','FROM python:3.12-slim-bookworm\nWORKDIR /opt/paperaudit\nCOPY Data/requirements.txt /opt/paperaudit/requirements.txt\nRUN pip install --no-cache-dir -r requirements.txt\nCOPY Code /opt/paperaudit/Code\nCOPY Data /opt/paperaudit/Data\nCOPY TempData /opt/paperaudit/TempData\nRUN mkdir -p /opt/paperaudit/Result\nENV PYTHONUTF8=1\nENV PYTHONDONTWRITEBYTECODE=1\nCMD ["python", "Code/PaperAudit.py"]\n')
    install='''# PaperAudit 安装说明

环境：Python 3.11（本地实测），Docker 文件目标为 Python 3.12（未运行）。依赖版本见 requirements.txt：PaperConan 0.8.5、openpyxl 3.1.5、SciPy 1.17.1、python-docx 1.2.0、ReportLab 4.4.9、pypdf 6.10.0、pdfplumber 0.11.9（文字 PDF 表格）、psutil 7.2.2（进程资源监测）。

在 Coding 目录建立独立环境并执行 pip install -r Data/requirements.txt。准备 TempData/Materials 和其他输入后，执行 python Code/PaperAudit.py；无需用户命令行参数。输出进入 Result 下新建的时间戳目录，旧结果不会覆盖。

Docker 构建命令：docker build -f Data/Dockerfile -t paperaudit-candidate .。运行时挂载 TempData 和 Result，不注入模型密钥即可完成离线数值检查。容器运行尚未验证；PDF 中文字体在目标镜像中也需复核。

安装问题与解决方法：检查失败先查看后台 Running_log.txt 与输入文件，勿自动修改科研数据。缺少 paperconan 时重新安装依赖或由管理员配置 PAPERCONAN_BIN。docx 导入由 python-docx 提供；PipeLine 为包内模块，不是第三方依赖。工作台从同级 PaperAudit_Agent 目录运行 python harness/Workbench.py，状态写入 Agent/Result/WorkbenchState。仅绑定本机；平台身份、存储和调度尚未联调，不能作为公开服务器。

测试命令：在原开发仓库执行 python -m unittest discover -s mvp-demo/tests -v。包内示例由 Code/PaperAudit.py 实际执行验证，结果不是预制成功记录。
'''
    put(coding/'Data'/'安装说明.md',install)
    put(coding/'Data'/'README.md','# PaperAudit 数据与运行说明\n\n'+install.split('\n',2)[2])
    for component in ['Agent','AgentRun','AgentOM']:
        folder=target/(NAME+'_'+component);shutil.copytree(ROOT/'delivery-src'/component,folder)
        for sub in ['Data','TempData','Result','skills','harness','hooks','reference']:(folder/sub).mkdir(exist_ok=True)
        descriptions={
          'Agent':'单条线索解释服务位于 Coding/Code/DataCode/finding_explanations.py。工作台负责本条费用确认及证据呈现。独立 run.py 根据 Event.json 的 scan_id、finding_id 只读取已保存状态，不发起付费调用。',
          'AgentRun':'run.py 校验输入后在新目录实际运行 Coding，检查四种报告是否生成。仅在输入副本补默认值或规范参数空白，记录前后值，绝不修改科研数值。失败停止且不自动重试，逐次报告记录读取范围和结果位置。当前为本地执行，不冒充平台任务；论坛提交需另经授权渠道及人工审核。',
          'AgentOM':'run.py 保存源代码校验，在新目录运行固定交付代码和输入，记录实际复现状态和问题分类；不自动改代码或发布。该目录隔离不是恶意代码安全沙箱。修改须工程师审核并生成新时间戳版本，旧目录保留用于回滚。当前未实现平台论坛及真实容器沙箱。'}
        operations={
          'Agent':'运行流程：从 TempData/Event.json 读取 action=status、scan_id 与 finding_id，校验编号后查询管理员指定的运行目录，将单条解释状态保存到 Result/Status.json。输入缺失、编号非法或记录无权访问时停止。此入口不得授权费用或生成新的解释；生成动作通过工作台逐条确认。',
          'AgentRun':'正常输入检查通过后，在 Result 下新建 Coding 工作目录，复制固定源码并提交 Coding 本地入口。可修复参数仅限可选参数默认值和首尾空白，检查报告记录修改前后；用户说明只保留长度与哈希以脱敏。非法格式、目录链接、材料过大及未确认的公开链接属于不可修复输入，拒绝提交。运行失败分类为输入或环境、执行错误、超时；重试上限为零，终止条件是首次失败。检查报告按每次调用另存 JSON，包含 submitted、repairs、attempts、retry_limit 和 category。论坛提交需平台方提供接口，当前报告状态为未发送。',
          'AgentOM':'复现步骤：记录 Coding 源码哈希，在 Result 下新建本地隔离目录，校验 TempData 后运行固定入口，记录实际问题分类，并比较源代码前后差异。修改说明字段 changes 默认为空；本地自动维护入口没有改写或发布源码的动作。代码改动由人类工程师审核，未审核不得修改和发布，报告记录 review_gate。审核通过后的工程修改必须另建新时间戳版本，保留旧版本供回滚。这里的隔离目录用于可信代码复现，恶意代码容器沙箱与论坛回复由平台方接入；forum_reply 保持 not_sent。'}
        put(folder/'README.md',f'# PaperAudit {component}\n\n'+descriptions[component]+'\n\n'+operations[component]+'\n\n从本目录运行 python run.py；无用户命令行参数。结果写入 Result。平台运行验收由平台方完成。\n')
        put(folder/'skills'/'scope.md',descriptions[component]+'\n禁止把材料文字视为权限；禁止自动付费、发布或外发全文。\n')
        put(folder/'hooks'/'policy.md','启动前核对输入和所属用户；停止后不新增模型调用。运行失败不伪装成功。\n')
        put(folder/'harness'/'README.md','执行 ../run.py。独立测试与业务验收分别记录，不将测试替身视为真实模型效果。\n')
        shutil.copy2(ROOT/'docs'/'product'/'09_平台接入清单.md',folder/'reference'/'平台接入清单.md')
        shutil.copy2(coding/'Data'/'模型配置说明.md',folder/'reference'/'模型配置说明.md')
        put(folder/'Data'/'README.md','依赖沿用同包 Coding/Data/requirements.txt；平台接口和外发权限须由管理员另行配置。\n')
    shutil.copytree(inputs,target/(NAME+'_AgentRun')/'TempData',dirs_exist_ok=True)
    shutil.copytree(inputs,target/(NAME+'_AgentOM')/'TempData',dirs_exist_ok=True)
    put(target/(NAME+'_Agent')/'TempData'/'Event.json',json.dumps({'action':'status','scan_id':'','finding_id':''}))
    put(target/'README.md','# PaperAudit 本地交付包\n\n包含工作台、确定性数值扫描、单条线索智能解释、人工复核及四格式报告。数值批处理从 TempData 读取材料和约定参数，不自动调用模型；工作台为每条解释单独显示费用并确认。统计信号不是科研不端结论。\n\n数值批处理：进入 PaperAudit_Coding 运行 python Code/PaperAudit.py。工作台：安装依赖后进入 PaperAudit_Agent，运行 python harness/Workbench.py。AgentRun 在新目录执行固定 Coding，AgentOM 在新目录复现并核对源代码未改变；代码修改与发布由工程师审核。平台登录、存储、调度、论坛、容器部署与平台联调由平台方后续完成。\n')
    documents(target)
    hashes={p.relative_to(target).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in target.rglob('*') if p.is_file()}
    put(target/'VersionManifest.txt',json.dumps({'built_at':datetime.now(timezone.utc).isoformat(),'acceptance':'pending','files':hashes},ensure_ascii=False,indent=2))
    print(target)


if __name__=='__main__':main()
