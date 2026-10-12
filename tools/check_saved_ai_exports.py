"""Read all four actual exports containing retained REAL replies, no model call."""
import html
import json
import re
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'test-results/2026-10-11-multipaper-v1'
def normalized(s):return re.sub(r'\s+','',s)
records=[]
for record in json.loads((OUT/'live-ai-summary.json').read_text(encoding='utf-8'))['records']:
    folder=OUT/record['case_id']/'after-ai';expected=record['state']['explanation']
    doc=Document(folder/'report.docx');word='\n'.join(p.text for p in doc.paragraphs)
    pdf='\n'.join(p.extract_text() or '' for p in PdfReader(folder/'report.pdf').pages)
    book=load_workbook(folder/'details.xlsx',read_only=True);excel='\n'.join(str(v) for s in book for row in s.values for v in row);book.close()
    web=html.unescape((folder/'report.html').read_text(encoding='utf-8'))
    checked=[]
    for name,content in [('Word',word),('PDF',pdf),('Excel',excel),('HTML',web)]:
        for text in [expected['what'],expected['why'],expected['report_summary'],expected['disclaimer'],*expected['benign_explanations'],*expected['check_steps']]:
            assert normalized(text) in normalized(content),f'AI field lost: {record["case_id"]}/{name}'
        checked.append(name)
    records.append({'case_id':record['case_id'],'passed':True,'formats':checked,'model_calls_added':0})
(OUT/'real-ai-all-formats.json').write_text(json.dumps({'records':records,'passed':len(records),'failed':0,'model_calls_added':0},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':len(records),'failed':0,'model_calls_added':0}))
