"""Read actual real-paper exports independently; do not manufacture review decisions."""
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
folder=ROOT/'docs/verification/2026-09-30-real-paper'
browser=json.loads((folder/'browser.json').read_text(encoding='utf-8'))
scan=browser['scan_id']
runtime=folder/'workbench-runtime'/scan
result=json.loads((runtime/'result.json').read_text(encoding='utf-8'))
reports=sorted((runtime/'reports').iterdir(),key=lambda p:p.stat().st_mtime)
report=reports[-1]
doc=Document(report/'report.docx')
word='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
pdf=PdfReader(report/'report.pdf')
pages=[p.extract_text() or '' for p in pdf.pages]
assert all(len(text.strip())>30 for text in pages),'Blank PDF page'
book=load_workbook(report/'details.xlsx',read_only=True,data_only=False)
try:
    cells=[c for sheet in book for row in sheet for c in row if c.value is not None]
    assert all(c.data_type!='f' for c in cells),'Unexpected executable formula'
    excel='\n'.join(str(c.value) for c in cells)
finally:book.close()
texts={'Word':word,'PDF':'\n'.join(pages),'Excel':excel,'HTML':(report/'report.html').read_text(encoding='utf-8')}
values=[str(c['original']) for f in result['findings'] for c in f['evidence'].get('focus_cells',[])]
assert len(values)==218
for name,text in texts.items():
    assert all(value in text for value in values),name+' lost raw precision'
    assert report.name in text,name+' missing same snapshot id'
    assert '保留1位小数后' in text,name+' stale wording'
    assert '论文内容核对尚未开通' in text,name+' missing scope'
record={'at':datetime.now(timezone.utc).isoformat(),'passed':True,'scan_id':scan,'report_id':report.name,
        'formats_checked':list(texts),'raw_cell_values_checked_per_format':len(values),'pdf_pages':len(pages),
        'hashes':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in report.iterdir() if p.is_file()},
        'boundary':'Full content checks, not independent research conclusions or a claim of per-page visual PDF acceptance.'}
(folder/'reports.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(record,ensure_ascii=False))
