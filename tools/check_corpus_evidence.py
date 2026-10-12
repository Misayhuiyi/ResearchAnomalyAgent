"""Independent source-value/report checks; known errors are NOT inferred labels."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
CORPUS=ROOT/'test-corpus'


def read_json(path):return json.loads(path.read_text(encoding='utf-8'))


def table(path,sheet):
    if path.suffix.lower() in {'.csv','.tsv'}:
        with path.open(encoding='utf-8-sig',newline='') as handle:
            return list(csv.reader(handle,delimiter='\t' if path.suffix=='.tsv' else ','))
    if path.suffix.lower()=='.xlsx':
        book=load_workbook(path,read_only=True,data_only=True)
        try:return [list(row) for row in book[sheet].values]
        finally:book.close()
    return None


def same(a,b):
    if a==b or str(a)==str(b):return True
    try:return math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-10)
    except (ValueError,TypeError):return a in ('',None) and b in ('',None)


def numeric(value):
    if value in (None,'') or isinstance(value,bool):return False
    try:return math.isfinite(float(value))
    except (TypeError,ValueError):return False


def check_case(case,folder):
    result_path=folder/'result.json'
    if folder.name=='after-ai' and not result_path.exists():result_path=folder.parent/'result.json'
    r=read_json(result_path); cache={}; checked=0; numeric_checked=0; mismatches=[]; unaddressable=[]; addressed=set()
    source={Path(p).name:CORPUS/p for p in case['uploads']}
    if 'data_rows' in case['expected']:
        path=next(p for p in source.values() if p.suffix in ('.csv','.tsv'))
        original=table(path,path.stem)
        assert len(original)-1==case['expected']['data_rows'],'Data row count differs from source contract'
        if 'groups' in case['expected']:
            assert len({row[0] for row in original[1:]})==case['expected']['groups'],'Group count differs'
    def get(file,sheet):
        key=file,sheet
        if key not in cache:cache[key]=table(source[file],sheet) if file in source else None
        return cache[key]
    def validate(file,sheet,row,column,value,fid):
        nonlocal checked,numeric_checked
        rows=get(file,sheet)
        if rows is None:return  # PDF evidence has a separate extraction boundary.
        if not isinstance(row,int) or not 1<=row<=len(rows):
            mismatches.append({'finding':fid,'reason':'row outside source','row':row});return
        if isinstance(column,int):idx=column-1
        else:
            # Real books may have title rows; match a UNIQUE header in the first 3 rows.
            matches={i for headers in rows[:3] for i,h in enumerate(headers) if h not in (None,'') and str(h)==str(column)}
            fallback=re.fullmatch(r'第 (\d+) 列',str(column))
            if len(matches)==1:idx=next(iter(matches))
            elif fallback:idx=int(fallback[1])-1
            else:
                unaddressable.append({'finding':fid,'file':file,'sheet':sheet,'row':row,'column':column});return
        checked+=1;numeric_checked+=int(numeric(value));addressed.add((file,sheet,row,idx+1))
        if idx>=len(rows[row-1]) or not same(rows[row-1][idx],value):
            mismatches.append({'finding':fid,'file':file,'sheet':sheet,'row':row,'column':idx+1,'shown':value,'source':rows[row-1][idx] if idx<len(rows[row-1]) else None})
    missing_rounding=[]
    for f in r['findings']:
        assert all(f.get(k) for k in ('id','title_cn','summary_cn','why_cn','next_steps')), 'Chinese explanation missing'
        ev=f.get('evidence') or {}
        for cell in ev.get('focus_cells',[]):
            validate(f['file'],f['sheet'],cell.get('row_idx'),cell.get('column'),cell.get('original'),f['id'])
        for pair in ev.get('comparison',[]):
            for side in ('left','right'):
                cell=pair[side]
                validate(cell.get('file',f['file']),cell.get('sheet',f['sheet']),cell.get('row_idx'),cell.get('column'),cell.get('value'),f['id'])
        if f['kind']=='identical_after_rounding' and len(ev.get('focus_cells',[]))!=f['n']:
            missing_rounding.append({'finding':f['id'],'expected':f['n'],'available':len(ev.get('focus_cells',[]))})
        rows=get(f['file'],f['sheet'])
        if rows is not None:
            # Compare visible evidence slices against the original row, independently of column labels.
            for shown in ev.get('rows',[]):
                ri=shown.get('row_idx');values=shown.get('values') or []
                if not isinstance(ri,int) or not 1<=ri<=len(rows):
                    mismatches.append({'finding':f['id'],'reason':'preview row outside source'});continue
                actual=rows[ri-1]
                offsets=[start for start in range(len(actual)-len(values)+1) if all(same(a,b) for a,b in zip(actual[start:start+len(values)],values))]
                if not offsets:mismatches.append({'finding':f['id'],'reason':'preview differs from source','row':ri})
                else:
                    checked+=len(values);numeric_checked+=sum(numeric(v) for v in values)
    # Determine only direct evidence coverage of documented cells; not semantic error detection.
    known=[]
    for cell in case['expected'].get('known_cells',[]):
        sheet=cell.get('sheet',Path(cell['file']).stem)
        rows=get(cell['file'],sheet);assert same(rows[cell['row']-1][cell['column']-1],cell['original'])
        known.append({**cell,'highlighted':(cell['file'],sheet,cell['row'],cell['column']) in addressed,
                      'error_diagnosed':False,'reason':'Highlight is only a statistical signal; comparison truth comes from source correction.'})
    for cell in case['expected'].get('corrected_cells',[]):
        rows=get(cell['file'],cell.get('sheet',Path(cell['file']).stem))
        assert same(rows[cell['row']-1][cell['column']-1],cell['value'])
    # Each actual export must be readable, not just have a successful HTTP response.
    doc=Document(folder/'report.docx');word='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    reader=PdfReader(folder/'report.pdf');pdf='\n'.join(p.extract_text() or '' for p in reader.pages)
    book=load_workbook(folder/'details.xlsx',read_only=True,data_only=False);excel=[];formula_count=0
    metadata=dict(book['报告概览'].values)
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type=='f':formula_count+=1
                if cell.value is not None:excel.append(str(cell.value))
    book.close();excel='\n'.join(excel)
    html=(folder/'report.html').read_text(encoding='utf-8')
    # Reports identify their immutable snapshot, not the scan's task ID.
    report_id=metadata.get('报告编号')
    assert report_id and all(report_id in s for s in (word,pdf,excel,html)), 'Export report identity differs'
    source_hash=hashlib.sha256(json.dumps(r,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
    assert metadata.get('原检查结果校验')==source_hash, 'Export source differs from scanned result'
    assert all(source_hash in s.replace('\n','') for s in (word,pdf,excel,html)), 'Export source binding missing'
    assert formula_count==0,'User text became Excel formula'
    precision=[]
    for f in r['findings']:
        if f['file'].endswith(('.csv','.tsv')):
            for cell in f.get('evidence',{}).get('focus_cells',[]):
                token=str(cell['original'])
                precision.append({'value':token,'all_four_formats':all(token in s for s in (word,pdf,excel,html))})
    assert all(p['all_four_formats'] for p in precision),'Full original value lost in exported report'
    partial=[{'filename':m['filename'],'status':m['status']} for m in r['materials'] if m['status']!='read']
    return {'case_id':case['id'],'scan_id':r['scan_id'],'kinds':dict(Counter(f['kind'] for f in r['findings'])),
            'source_values_checked':checked,'source_mismatches':mismatches,'unaddressable_cells':unaddressable,
            'numeric_values_checked':numeric_checked,
            'missing_rounding_evidence':missing_rounding,'known_cells':known,'pdf_pages':len(reader.pages),
            'csv_full_precision_exports_checked':len(precision),'formula_count':formula_count,
            'partial_materials':partial,'coverage_notes':r.get('coverage_notes_cn',[]),
            'passed':not mismatches,'scope':'addressable numeric evidence and export integrity, not error truth or human acceptance'}


def main():
    parser=argparse.ArgumentParser();parser.add_argument('results');args=parser.parse_args()
    out=Path(args.results).resolve()
    if not out.is_relative_to(ROOT/'test-results'):raise ValueError('Local test-results required')
    manifest=read_json(CORPUS/'manifest.json');records=[]
    for case in manifest['cases']:
        folder=out/case['id']
        try:
            value=check_case(case,folder)
            if (folder/'after-ai/report.html').exists():
                value['after_ai_exports']=check_case(case,folder/'after-ai')
                value['passed']=value['passed'] and value['after_ai_exports']['passed']
        except Exception as exc:value={'case_id':case['id'],'passed':False,'error':str(exc)}
        records.append(value);(folder/'source-and-report-check.json').write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps({k:value.get(k) for k in ('case_id','passed','source_values_checked','error')},ensure_ascii=False),flush=True)
    summary={'records':records,'passed':sum(r['passed'] for r in records),'failed':sum(not r['passed'] for r in records)}
    (out/'evidence-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if all(r['passed'] for r in records) else 1


if __name__=='__main__':raise SystemExit(main())
