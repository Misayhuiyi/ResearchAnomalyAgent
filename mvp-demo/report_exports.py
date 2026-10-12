"""Four formats rendered from one immutable review snapshot, published as a set."""
from __future__ import annotations

import hashlib
import html
import json
import re
import threading
from pathlib import Path

from review_store import ReviewError, Unavailable, STATES

EXPORT_LOCK = threading.Lock()
FILES = {'report.html', 'report.docx', 'report.pdf', 'details.xlsx'}
MAX_EVIDENCE = 50000
MAX_TEXT = 4_000_000


def text(value):
    # XML 1.0 forbids control characters; do not let untrusted input break all formats.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value if value is not None else ''))


def evidence_rows(finding):
    ev = finding.get('evidence') or {}
    base = [finding.get('id'), finding.get('title_cn')]
    if ev.get('comparison'):
        for pair in ev['comparison']:
            for side, label in [('left', '第一组'), ('right', '第二组')]:
                cell = pair[side]
                yield base + [label, cell.get('file'), cell.get('sheet'), cell.get('row_idx'), cell.get('column'), cell.get('value'), '']
    elif ev.get('focus_cells'):
        for cell in ev['focus_cells']:
            yield base + ['', finding.get('file'), finding.get('sheet'), cell.get('row_idx'), cell.get('column'), cell.get('original'), cell.get('displayed')]
    else:
        headers = ev.get('headers') or []
        for row in ev.get('rows') or []:
            for index, value in enumerate(row.get('values') or []):
                yield base + ['', finding.get('file'), finding.get('sheet'), row.get('row_idx'), headers[index] if index < len(headers) else str(index+1), value, '']


def sections(snapshot):
    result, review = snapshot['result'], snapshot['review']
    reviews = review['reviews']
    findings = result.get('findings', [])
    pending = sum(reviews.get(f['id'], {}).get('state', 'pending') == 'pending' for f in findings)
    intro = [
        ['材料', result.get('source', '')], ['报告编号', snapshot['report_id']],
        ['材料版本',str(result.get('material_version',1))],
        ['复核版本', str(snapshot['revision'])], ['生成时间', snapshot['created_at']],
        ['内容校验', snapshot['snapshot_hash']], ['原检查结果校验', snapshot['source_hash']],
        ['检查范围', '仅包含本次已读取且实际执行的项目，不代表全部需求均已检查。'],
        ['材料读取', {'complete':'本次读取完成','partial':'部分内容未完成','failed':'未完成'}.get(result.get('scan_status'), '范围待确认')],
        ['论文内容核对', (snapshot.get('content_review') or {}).get('message') or '尚未执行论文内容核对；正文、主表与补充材料的含义对应仍需人工确认。'],
        ['待复核', str(pending)], ['使用边界', '结果仅供辅助核对，不能单独认定科研不端。'],
    ]
    for note in result.get('coverage_notes_cn', []): intro.append(['范围说明', note])
    output = [('报告概览', ['项目', '内容'], intro)]
    materials = [[m.get('filename'), m.get('source_url') or '用户上传', m.get('sha256'), m.get('size')]
                 for m in result.get('material_sources', [])]
    output.append(('材料来源', ['材料', '来源', '内容校验', '字节数'], materials))
    output.append(('材料读取', ['材料', '读取情况', '说明'], [
        [m.get('filename'), {'read':'已读取','partial':'部分读取','failed':'读取失败','unread':'未读取'}.get(m.get('status'),'未确认'),m.get('reason_cn','')]
        for m in result.get('materials', [])]))
    output.append(('核对范围', ['材料', '数据表', '核对项目', '状态', '计算口径', '未完成原因'], [
        [c.get('file'), c.get('sheet'), c.get('name'), {'checked':'已核对','partial':'部分核对','insufficient':'条件不足','unsupported':'暂不支持','failed':'未完成'}.get(c.get('status'), '未确认'),
         c.get('assumption'), '；'.join(c.get('notes', []))] for c in result.get('statistical_coverage', [])]))
    found, evidence = [], []
    for f in findings:
        saved = reviews.get(f['id'], {'state': 'pending', 'note': ''})
        found.append([f['id'], f.get('title_cn'), f.get('location'), f.get('summary_cn'),
                      f.get('rule',''), '；'.join(m.get('summary','') for m in f.get('matched_rules', [])),
                      STATES[saved['state']], saved['note'], (f.get('evidence') or {}).get('note_cn', '')])
        evidence.extend(evidence_rows(f))
        if len(evidence) > MAX_EVIDENCE:
            raise ReviewError('报告原始值超过导出容量，请拆分材料；本次未生成完整报告。')
    output.append(('核对项目', ['编号', '项目', '位置', '核对结果', '计算关系', '同位置的其他关系', '复核状态', '复核依据', '原始值范围'], found))
    output.append(('原始值', ['编号', '项目', '组别', '文件', '数据表', '行号', '列名', '原始值', '保留1位小数后'], evidence))
    content = snapshot.get('content_review') or {}
    if content.get('budget_id'):
        output.append(('内容核对范围', ['项目','内容'], [
            ['范围说明',content.get('boundary','')], ['计费币种',content.get('currency')],
            ['费用口径','费用按配置单价保守估算，实际扣费以 DeepSeek 账单为准。'],
            ['累计授权上限',str(content.get('cap_micros',0)/1000000)],
            ['已记录费用',str(content.get('spent_micros',0)/1000000)],
            ['待核对费用',str(content.get('held_micros',0)/1000000)],
            ['本论文各版本累计已记录费用',str(content.get('paper_spent_micros',0)/1000000)],
            ['完成批次',f"{content.get('next_step',0)} / {content.get('planned_calls',0)}"],
            *[['材料范围',f"{c['file']} · {c['source_count']} 个片段 · {c.get('reason','')}"] for c in content.get('coverage',[])]]))
        suggestions, comparisons = [], []
        for step in content.get('outputs',[]):
            for item in step['suggestions']:
                for citation in item['citations']:
                    suggestions.append([item['finding_id'],item['summary'],citation['file'],citation['location'],citation['quote'],'仅为建议，未改变人工复核'])
            for item in step['tool_results']:
                if 'summary' not in item: continue
                for citation in item.get('citations',[]) or [{}]:
                    comparisons.append([item['summary'],citation.get('file',''),citation.get('location',''),citation.get('original',''),citation.get('unit',''),'对应关系需人工确认'])
        output.append(('原文核对建议', ['核对编号','建议','材料','位置','原文','说明'],suggestions))
        output.append(('跨材料数值核对', ['核对结果','材料','位置','原始值','单位','说明'],comparisons))
    output.append(('复核记录', ['版本', '时间', '核对编号', '操作人', '原状态', '新状态', '说明'], [
        [e['revision'], e['at'], e['finding'], e['operator'], STATES[e['before']['state']],
         STATES[e['after']['state']], e['after']['note']] for e in snapshot.get('history', [])]))
    total = sum(len(text(c)) for _, headers, rows in output for row in [headers, *rows] for c in row)
    if total > MAX_TEXT:
        raise ReviewError('报告内容超过导出容量，请拆分材料；本次未生成完整报告。')
    return [(title, list(map(text, heads)), [list(map(text, row)) for row in rows]) for title, heads, rows in output]


def write_html(path, tables):
    parts = ['<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>科研数据核对报告</title>',
             '<style>body{font:15px/1.7 "Microsoft YaHei",sans-serif;max-width:1100px;margin:35px auto;color:#18344b;padding:20px}table{border-collapse:collapse;width:100%;table-layout:fixed}td,th{border:1px solid #d8e2ed;padding:7px;word-break:break-all;text-align:left;vertical-align:top}th{background:#edf4fc}h2{margin-top:35px}@media print{thead{display:table-header-group}tr{break-inside:avoid}}</style><h1>科研数据核对报告</h1>']
    for title, headers, rows in tables:
        parts.append('<h2>'+html.escape(title)+'</h2><table><thead><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in headers)+'</tr></thead><tbody>')
        for row in rows:
            parts.append('<tr>'+''.join('<td>'+html.escape(c).replace('\n','<br>')+'</td>' for c in row)+'</tr>')
        parts.append('</tbody></table>' if rows else '</tbody></table><p>无记录。</p>')
    parts.append('</html>')
    path.write_text(''.join(parts), encoding='utf-8')


def write_docx(path, tables):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.7)
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for name in ['Normal','Title','Heading 1']:
        selected=doc.styles[name]
        selected.font.color.rgb=RGBColor.from_string('000000')
        selected.font.underline=False
        ppr=selected.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
    doc.add_heading('科研数据核对报告', 0)
    # Wide records use vertical labelled entries instead of tiny nine-column text.
    for title, headers, rows in tables:
        doc.add_heading(title, 1)
        if not rows: doc.add_paragraph('无记录。')
        elif len(headers) <= 2:
            for row in rows:
                paragraph=doc.add_paragraph()
                paragraph.paragraph_format.space_after=Pt(3)
                paragraph.add_run(row[0]+'：').bold=True
                paragraph.add_run(row[1])
        else:
            for i, row in enumerate(rows, 1):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.keep_together=True
                paragraph.paragraph_format.space_after=Pt(8)
                paragraph.add_run(f'{i}. ').bold = True
                for j, (label, value) in enumerate(zip(headers, row)):
                    if not value:continue
                    if j: paragraph.add_run('；')
                    paragraph.add_run(label+'：').bold = True
                    paragraph.add_run(value)
    footer = section.footer.paragraphs[0]
    footer.add_run('科研数据核对报告 · 第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    doc.save(path)


def write_pdf(path, tables):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    normal = ParagraphStyle('CN', fontName='STSong-Light', fontSize=10, leading=15, wordWrap='CJK', spaceAfter=5)
    heading = ParagraphStyle('CNHead', parent=normal, fontSize=15, leading=22, spaceBefore=14, spaceAfter=8, keepWithNext=True)
    story = [Paragraph('科研数据核对报告', heading)]
    for title, headers, rows in tables:
        story.append(Paragraph(html.escape(title), heading))
        if not rows: story.append(Paragraph('无记录。', normal))
        for row in rows:
            record=(row[0]+'：'+row[1]) if len(headers)==2 else '；'.join(label+'：'+value for label,value in zip(headers,row) if value)
            paragraph=Paragraph(html.escape(record).replace('\n','<br/>'),normal)
            # Keep short records intact; long notes remain splittable across pages.
            story.append(KeepTogether([paragraph]) if len(record)<500 else paragraph)
            story.append(Spacer(1, 3))
    def footer(canvas, doc):
        canvas.setFont('STSong-Light', 9)
        canvas.drawRightString(A4[0]-36, 20, f'科研数据核对报告 · 第 {doc.page} 页')
    SimpleDocTemplate(str(path), pagesize=A4, topMargin=30, bottomMargin=36, leftMargin=36,
                      rightMargin=36, title='科研数据核对报告').build(story, onFirstPage=footer, onLaterPages=footer)


def write_xlsx(path, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook(); book.remove(book.active)
    for title, headers, rows in tables:
        sheet = book.create_sheet(title)
        for row in [headers, *rows]:
            if any(len(value) > 32767 for value in row):
                raise ReviewError('单个明细超出 Excel 单元格容量，未生成报告。')
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.data_type = 's'  # Never execute a user-supplied formula or URL.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='214E76')
        from openpyxl.utils import get_column_letter
        for index in range(1,len(headers)+1): sheet.column_dimensions[get_column_letter(index)].width = 26
    book.save(path)


def export_snapshot(store, actor, snapshot):
    """Retryable same-snapshot export. No download is published until every file is ready."""
    with EXPORT_LOCK:
        current = store.snapshot(actor, snapshot['report_id'])
        from filesystem_paths import filesystem_path
        folder = filesystem_path(store.root / current['scan_id'] / 'reports' / current['report_id'])
        if current['ready']:
            for name in FILES: download_path(store, actor, current['report_id'], name)
            return export_links(current)
        tables = sections(current)
        folder.mkdir(parents=True, exist_ok=True)
        for name, writer in [('report.html',write_html),('report.docx',write_docx),('report.pdf',write_pdf),('details.xlsx',write_xlsx)]:
            writer(folder / name, tables)
        manifest = {'snapshot_hash': current['snapshot_hash'], 'files': {
            name: hashlib.sha256((folder/name).read_bytes()).hexdigest() for name in sorted(FILES)}}
        pending = folder/'manifest.pending'
        pending.write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
        pending.replace(folder/'manifest.json')
        store.mark_ready(actor, current['report_id'])
        return export_links(current)


def export_links(snapshot):
    return {'report_id': snapshot['report_id'], 'revision': snapshot['revision'],
            'snapshot_hash': snapshot['snapshot_hash'], 'downloads': {
                name: f"/api/reports/{snapshot['report_id']}/{name}" for name in sorted(FILES)}}


def download_path(store, actor, report_id, filename):
    if filename not in FILES: raise Unavailable('报告文件不存在。')
    snap = store.snapshot(actor, report_id)
    if not snap['ready']: raise ReviewError('报告尚未完整生成，请重试生成。')
    from filesystem_paths import filesystem_path
    folder = filesystem_path(store.root/snap['scan_id']/'reports'/report_id)
    try:
        manifest = json.loads((folder/'manifest.json').read_text(encoding='utf-8'))
        path = folder/filename
        if manifest['snapshot_hash'] != snap['snapshot_hash'] or hashlib.sha256(path.read_bytes()).hexdigest() != manifest['files'][filename]:
            raise ValueError('mismatch')
        return path
    except (OSError, ValueError, KeyError):
        raise ReviewError('报告文件校验未通过，暂不能下载。') from None
