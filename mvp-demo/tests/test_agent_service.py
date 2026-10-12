import hashlib
import json
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import Mock
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from agent_service import AgentService
from budget_store import MODEL, VERSION, BudgetError
from model_gateway import ENDPOINT, ModelError
from platform_ports import LOCAL_ACTOR, Actor
from review_store import ReviewStore, ReviewError
from report_exports import sections
from source_catalog import build
import job_store

PRICE={'model':MODEL,'version':VERSION,'currency':'CNY','unit':'per_million_tokens','input':'2','output':'3','reference':'SYNTHETIC TEST ONLY'}


class AgentTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        self.scan='abcdef123456';self.input=self.root/self.scan/'input';self.input.mkdir(parents=True)
        self.data=b'group,value\ncontrol,1 mg\ntreated,2 mg\n';(self.input/'table.csv').write_bytes(self.data)
        self.manifest=[{'filename':'table.csv','sha256':hashlib.sha256(self.data).hexdigest(),'size':len(self.data)}]
        self.store=ReviewStore(self.root);self.store.register(LOCAL_ACTOR,{'scan_id':self.scan,'material_sources':self.manifest,'findings':[{'id':'F1'}]})
        self.config={'endpoint':ENDPOINT,'price':PRICE,'secret':'TEST-ONLY','valid_until':time.time()+1000}
        self.transport=Mock(side_effect=self.answer)
        self.service=AgentService(self.root,lambda:self.config,self.transport)
        self.roles={'table.csv':'table'}

    def answer(self,secret,body):
        request=json.loads(body);sources=json.loads(request['messages'][1]['content'])['sources']
        sid,source=next(iter(sources.items()))
        content={'suggestions':[{'finding_id':'F1','reason':'needs_context','citations':[{'source_id':sid,'quote':source['text']}]}],'tool_requests':[]}
        return {'model':MODEL,'choices':[{'finish_reason':'stop','message':{'content':json.dumps(content)}}],'usage':{'prompt_tokens':100,'completion_tokens':100}}

    def prepare(self):return self.service.prepare(LOCAL_ACTOR,self.scan,self.roles)
    def start(self,prepared,confirmed=True):
        return self.service.start(LOCAL_ACTOR,self.scan,prepared['quote']['quote_id'],'confirm-0001',confirmed,background=False)

    def test_scope_quote_no_paid_io_until_explicit_consent(self):
        p=self.prepare();self.assertEqual(p['status'],'awaiting_confirmation');self.transport.assert_not_called()
        self.assertEqual(p['budget_id'],self.prepare()['budget_id'])
        with self.assertRaises(BudgetError):self.start(p,False)
        self.transport.assert_not_called()
        done=self.start(p);self.assertEqual(done['status'],'completed');self.assertEqual(len(done['outputs']),1)
        self.assertEqual(self.start(p)['status'],'completed');self.transport.assert_called_once()
        self.assertEqual(self.store.view(LOCAL_ACTOR,self.scan)['revision'],0)
        self.assertEqual(self.store.view(LOCAL_ACTOR,self.scan)['reviews'],{})

    def test_owner_scope_change_and_hash_tamper_refused(self):
        with self.assertRaises(ReviewError):self.service.view(Actor('other','someone'),self.scan)
        p=self.prepare()
        with self.assertRaises(ReviewError):self.service.prepare(LOCAL_ACTOR,self.scan,{'table.csv':'body'})
        (self.input/'table.csv').write_bytes(b'changed')
        with self.assertRaises(ValueError):build(self.input,self.manifest,self.roles)
        self.transport.assert_not_called()

    def test_uncertain_billing_and_restart_never_repeat(self):
        p=self.prepare();self.transport.side_effect=TimeoutError()
        state=self.start(p);self.assertEqual(state['status'],'uncertain');self.assertGreater(state['held_micros'],0)
        self.start(p);self.transport.assert_called_once()
        with self.assertRaises(ReviewError):self.service.quote(LOCAL_ACTOR,self.scan)

    def test_cancel_before_send_and_recover(self):
        p=self.prepare();self.service.cancel(LOCAL_ACTOR,self.scan)
        with self.assertRaises(BudgetError):self.start(p)
        self.transport.assert_not_called()

    def test_graceful_shutdown_prevents_next_paid_call(self):
        self.prepare()
        run=self.service._read(LOCAL_ACTOR,self.scan)
        run['status']='running';self.service._save(LOCAL_ACTOR,self.scan,run)
        self.service.stop_running(Actor('other','other'))
        self.assertEqual(self.service.view(LOCAL_ACTOR,self.scan)['status'],'running')
        self.service.stop_running(LOCAL_ACTOR)
        self.service.run(LOCAL_ACTOR,self.scan)
        self.assertEqual(self.service.view(LOCAL_ACTOR,self.scan)['status'],'cancelled')
        self.transport.assert_not_called()

    def test_suggestions_are_frozen_with_four_format_source(self):
        p=self.prepare();done=self.start(p)
        snap=self.store.freeze(LOCAL_ACTOR,self.scan,0,'freeze-test1',self.service.report_context(LOCAL_ACTOR,self.scan))
        tables=sections(snap);self.assertTrue(any(title=='原文核对建议' and rows for title,_,rows in tables))
        self.assertEqual(snap['content_review']['outputs'],done['outputs'])
        self.assertNotIn('secret',json.dumps(snap));self.assertNotIn('TEST-ONLY',json.dumps(snap))

    def test_invalid_model_tools_are_not_adopted(self):
        self.transport=Mock(return_value={'model':MODEL,'choices':[{'finish_reason':'stop','message':{'content':'{"suggestions":[],"tool_requests":[{"name":"shell","arguments":{}}]}'}}],'usage':{'prompt_tokens':100,'completion_tokens':100}})
        self.service.transport=self.transport
        done=self.start(self.prepare());self.assertEqual(done['status'],'failed');self.assertEqual(done['outputs'],[])
        self.assertGreater(done['spent_micros'],0)

    def test_large_source_is_reduced_to_one_small_private_excerpt_batch(self):
        data=('group,value\n'+''.join(f'group{i},'+('x'*900)+'\n' for i in range(20))).encode()
        (self.input/'table.csv').write_bytes(data)
        # A separate scan is required for different materials.
        self.scan='abcdef123457';folder=self.root/self.scan/'input';folder.mkdir(parents=True);(folder/'table.csv').write_bytes(data)
        self.store.register(LOCAL_ACTOR,{'scan_id':self.scan,'material_sources':[{'filename':'table.csv','sha256':hashlib.sha256(data).hexdigest()}],'findings':[{'id':'F1'}]})
        p=self.prepare();self.assertEqual(p['planned_calls'],1)
        planned=self.service._read(LOCAL_ACTOR,self.scan)
        self.assertLess(len(json.dumps(planned['catalog']).encode()),4000)
        self.assertLess(len(planned['catalog']),20)
        done=self.start(p);self.assertEqual(done['status'],'completed')
        self.assertEqual(self.transport.call_count,1)
        self.assertLessEqual(done['spent_micros'],done['cap_micros'])

    def test_no_configuration_remains_usable_for_manual_review(self):
        def missing():raise ModelError('not configured')
        service=AgentService(self.root,missing)
        self.assertFalse(service.view(LOCAL_ACTOR,self.scan)['available'])
        with self.assertRaises(ModelError):service.prepare(LOCAL_ACTOR,self.scan,self.roles)
        self.assertEqual(self.store.view(LOCAL_ACTOR,self.scan)['revision'],0)

    def test_word_and_csv_provenance(self):
        from docx import Document
        doc=Document();doc.add_paragraph('Metric A control week1 adults 10 mg')
        table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Group';table.cell(0,1).text='Value';table.cell(1,0).text='control';table.cell(1,1).text='11 mg'
        path=self.input/'paper.docx';doc.save(path)
        extracted=build(self.input,[{'filename':'paper.docx','sha256':hashlib.sha256(path.read_bytes()).hexdigest()}],{'paper.docx':'body'})
        self.assertEqual({s['kind'] for s in extracted['catalog'].values()},{'body','table'})
        self.assertTrue(any('11 mg' in s['text'] and '行 2' in s['location'] for s in extracted['catalog'].values()))

    def test_material_version_preserves_paper_fees_and_uncertain_blocks_new_plan(self):
        original=self.scan
        job_store.save(self.root,{'job_id':original,'created_at':'a','paper_id':original})
        done=self.start(self.prepare())
        self.scan='abcdef123458';folder=self.root/self.scan/'input';folder.mkdir(parents=True)
        (folder/'table.csv').write_bytes(self.data)
        self.store.register(LOCAL_ACTOR,{'scan_id':self.scan,'material_sources':self.manifest,'findings':[{'id':'F1'}]})
        job_store.save(self.root,{'job_id':self.scan,'created_at':'b','paper_id':original})
        new=self.prepare();self.assertEqual(new['paper_spent_micros'],done['spent_micros'])
        self.transport.side_effect=TimeoutError();self.start(new)
        third='abcdef123459';self.store.register(LOCAL_ACTOR,{'scan_id':third,'material_sources':self.manifest,'findings':[]})
        job_store.save(self.root,{'job_id':third,'created_at':'c','paper_id':original})
        with self.assertRaises(ReviewError):self.service.prepare(LOCAL_ACTOR,third,self.roles)

    def test_restart_during_dispatch_is_not_automatically_resumed(self):
        p=self.prepare();run=self.service._read(LOCAL_ACTOR,self.scan)
        run['status']='running';self.service._save(LOCAL_ACTOR,self.scan,run)
        self.service.recover(LOCAL_ACTOR)
        self.assertEqual(self.service.view(LOCAL_ACTOR,self.scan)['status'],'uncertain')
        self.transport.assert_not_called()


if __name__=='__main__':unittest.main()
