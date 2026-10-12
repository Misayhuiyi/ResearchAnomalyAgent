import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from platform_ports import LOCAL_ACTOR, Actor
from review_store import ReviewStore, ReviewError, Unavailable
import report_exports as exports


class ExportTests(unittest.TestCase):
    def test_model_fee_is_not_presented_as_provider_invoice(self):
        snapshot = dict(self.snap, content_review={'budget_id':'test', 'currency':'USD'})
        table = next(rows for title, _, rows in exports.sections(snapshot) if title == '内容核对范围')
        self.assertIn(['费用口径','费用按配置单价保守估算，实际扣费以 DeepSeek 账单为准。'], table)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.store = ReviewStore(Path(self.tmp.name))
        self.result = {'scan_id':'abc123abc123','source':'试验.csv','scan_status':'partial',
                       'findings':[{'id':'FD001','title_cn':'数值需要核对','summary_cn':'13/40 与 25.0% 不一致。',
                                    'location':'试验.csv · 第3行','file':'试验.csv','sheet':'试验',
                                    'evidence':{'headers':['人数','原始百分比'], 'rows':[{'row_idx':3,'values':['13','=1+1']}],
                                                'note_cn':'仅检查本表。'}}]}
        self.store.register(LOCAL_ACTOR, self.result)
        self.store.save(LOCAL_ACTOR,self.result['scan_id'],'FD001','explained','核对原表，分母应为52。',0,'save-0001')
        self.snap = self.store.freeze(LOCAL_ACTOR,self.result['scan_id'],1,'freeze-0001')

    def test_same_snapshot_four_formats_and_formula_safety(self):
        from docx import Document
        from openpyxl import load_workbook
        from pypdf import PdfReader
        links = exports.export_snapshot(self.store,LOCAL_ACTOR,self.snap)
        content = {}
        for name in exports.FILES:
            path = exports.download_path(self.store,LOCAL_ACTOR,self.snap['report_id'],name)
            if name.endswith('.html'): content[name] = path.read_text(encoding='utf-8')
            if name.endswith('.docx'):
                doc=Document(path)
                content[name]='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
            if name.endswith('.pdf'): content[name]='\n'.join(p.extract_text() for p in PdfReader(path).pages)
            if name.endswith('.xlsx'):
                book=load_workbook(path)
                try:
                    content[name]='\n'.join(str(c.value) for s in book for row in s for c in row if c.value is not None)
                    cells=[c for s in book for row in s for c in row if c.value=='=1+1']
                    self.assertEqual(len(cells),1); self.assertEqual(cells[0].data_type,'s')
                finally: book.close()
        for name, value in content.items():
            with self.subTest(name=name):
                for required in ['核对原表，分母应为52。',self.snap['report_id'],self.snap['snapshot_hash'],'=1+1','尚未执行论文内容核对']:
                    self.assertIn(required,value)
        self.assertEqual(links,exports.export_snapshot(self.store,LOCAL_ACTOR,self.snap))
        with self.assertRaises(Unavailable): exports.download_path(self.store,Actor('x','x'),self.snap['report_id'],'report.pdf')
        with self.assertRaises(Unavailable): exports.download_path(self.store,LOCAL_ACTOR,self.snap['report_id'],'../result.json')

    def test_failure_never_publishes_partial_and_retry_same_version(self):
        with patch.object(exports,'write_pdf',side_effect=OSError('disk')):
            with self.assertRaises(OSError): exports.export_snapshot(self.store,LOCAL_ACTOR,self.snap)
        self.assertFalse(self.store.snapshot(LOCAL_ACTOR,self.snap['report_id'])['ready'])
        with self.assertRaises(ReviewError): exports.download_path(self.store,LOCAL_ACTOR,self.snap['report_id'],'report.html')
        self.assertEqual(exports.export_snapshot(self.store,LOCAL_ACTOR,self.snap)['report_id'],self.snap['report_id'])

    def test_tamper_and_html_injection(self):
        tables=[('标题',['列'],[['<script>alert(1)</script>']])]
        path=Path(self.tmp.name)/'safe.html'; exports.write_html(path,tables)
        self.assertNotIn('<script>',path.read_text(encoding='utf-8'))
        exports.export_snapshot(self.store,LOCAL_ACTOR,self.snap)
        pdf=exports.download_path(self.store,LOCAL_ACTOR,self.snap['report_id'],'report.pdf')
        pdf.write_bytes(b'tampered')
        with self.assertRaises(ReviewError): exports.download_path(self.store,LOCAL_ACTOR,self.snap['report_id'],'report.pdf')

    def test_capacity_fails_closed(self):
        with patch.object(exports,'MAX_EVIDENCE',1):
            with self.assertRaises(ReviewError): exports.export_snapshot(self.store,LOCAL_ACTOR,self.snap)
        self.assertFalse(self.store.snapshot(LOCAL_ACTOR,self.snap['report_id'])['ready'])

    def test_nested_windows_length_reports_are_downloadable(self):
        import shutil
        from filesystem_paths import filesystem_path, display_path
        root=filesystem_path(Path(self.tmp.name)/('a'*80)/('b'*80))
        store=ReviewStore(root);store.register(LOCAL_ACTOR,self.result)
        self.addCleanup(shutil.rmtree,root)
        snap=store.freeze(LOCAL_ACTOR,self.result['scan_id'],0,'long-path-001')
        exports.export_snapshot(store,LOCAL_ACTOR,snap)
        for name in exports.FILES:
            path=exports.download_path(store,LOCAL_ACTOR,snap['report_id'],name)
            self.assertGreater(len(display_path(path)),260)
            self.assertGreater(path.stat().st_size,100)


if __name__=='__main__': unittest.main()
