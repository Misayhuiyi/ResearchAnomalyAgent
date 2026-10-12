"""Batch screening, exact evidence, and restart-safe history regressions."""
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import app  # noqa: E402
import job_store  # noqa: E402


def _table_bytes(delimiter=",", offset=0):
    """Nonlinear measurements: repeatable evidence, not an index or dose ladder."""
    values = [1.123, 4.782, 2.356, 6.914, 3.257, 8.641, 7.385, 9.126,
              5.473, 2.849, 6.237, 4.516, 8.932, 3.684, 7.159, 5.827]
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=delimiter)
    writer.writerow(["sample", "measurement"])
    for index, value in enumerate(values, 1):
        writer.writerow([f"S{index:02}", f"{value + offset:.3f}"])
    return output.getvalue().encode("utf-8")


def _wait_for_job(job, runtime):
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        current = app.get_scan_job(job["job_id"], runtime_dir=runtime)
        if current["status"] in {"succeeded", "failed"}:
            return current
        time.sleep(0.05)
    raise AssertionError("后台批量筛查未在 60 秒内结束")


def _fresh_process(runtime, job_id=None, restore=False):
    # A fresh interpreter has an empty JOBS dict, so this tests actual persisted history.
    code = """
import json, sys
from pathlib import Path
import app
root = Path(sys.argv[1])
if sys.argv[3] == 'restore':
    app.restore_scan_jobs(root)
output = {'jobs': app.list_scan_jobs(runtime_dir=root)}
if sys.argv[2]:
    output['job'] = app.get_scan_job(sys.argv[2], runtime_dir=root)
print(json.dumps(output, ensure_ascii=True))
"""
    completed = subprocess.run(
        [sys.executable, "-c", code, str(runtime), job_id or "", "restore" if restore else ""],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=30,
        env={**os.environ, "PAPER_AUDIT_RUNTIME_DIR": str(runtime)}, check=True,
    )
    return json.loads(completed.stdout)


class BatchScreeningTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.storage = tempfile.TemporaryDirectory()
        cls.runtime = Path(cls.storage.name)
        cls.result = app.scan_files(
            [("experiment.csv", _table_bytes()), ("supplement.csv", _table_bytes())],
            runtime_dir=cls.runtime,
        )

    @classmethod
    def tearDownClass(cls):
        cls.storage.cleanup()

    def test_joint_scan_keeps_every_material_and_coverage(self):
        result = self.result
        self.assertEqual(["experiment.csv", "supplement.csv"], result["sources"])
        self.assertEqual(2, result["summary"]["files"])
        materials = {item["filename"]: item for item in result["materials"]}
        self.assertEqual(set(result["sources"]), set(materials))
        for item in materials.values():
            self.assertEqual("read", item["status"], item)
            self.assertEqual(1, item["sheets"])
            self.assertEqual(16, item["numeric_cells"])
        self.assertEqual("complete", result["coverage_summary"]["status"])
        self.assertTrue(result["coverage_summary"]["message_cn"])

    def test_cross_file_result_has_exact_both_sides(self):
        findings = [item for item in self.result["findings"] if item["category"] == "cross_sheet"]
        self.assertTrue(findings, "联合扫描应实际检出跨文件重复")
        pairs = [pair for finding in findings for pair in finding["evidence"].get("comparison", [])]
        self.assertTrue(pairs, "跨文件发现必须给出左右文件和原始数值，不能只保留泛化说明")
        seen = set()
        for pair in pairs:
            left, right = pair["left"], pair["right"]
            self.assertEqual({"experiment.csv", "supplement.csv"}, {left["file"], right["file"]})
            for cell in (left, right):
                self.assertEqual("measurement", cell["column"])
                self.assertGreaterEqual(cell["row_idx"], 2)
                self.assertEqual(Path(cell["file"]).stem, cell["sheet"])
            self.assertEqual(left["value"], right["value"])
            seen.add((left["row_idx"], float(left["value"]), right["row_idx"], float(right["value"])))
        self.assertIn((2, 1.123, 2, 1.123), seen)
        self.assertIn((3, 4.782, 3, 4.782), seen)

    def test_report_displays_material_coverage_and_cross_file_values(self):
        report = (self.runtime / self.result["scan_id"] / "report.html").read_text(encoding="utf-8")
        self.assertIn("experiment.csv", report)
        self.assertIn("supplement.csv", report)
        self.assertIn(self.result["coverage_summary"]["message_cn"], report)
        self.assertIn("1.123", report)
        self.assertIn("4.782", report)
        self.assertIn("measurement", report)
        self.assertNotIn("未发现需要核对的数据。", report)

    def test_same_sheet_name_in_different_files_cannot_mix_original_values(self):
        sample = (ROOT / "samples" / "fixed_shift_demo.csv").read_text(encoding="utf-8-sig")
        records = list(csv.reader(io.StringIO(sample)))
        reordered = io.StringIO(newline="")
        writer = csv.writer(reordered, delimiter="\t")
        writer.writerow(records[0])
        writer.writerows(reversed(records[1:]))
        with tempfile.TemporaryDirectory() as folder:
            result = app.scan_files(
                [("same.csv", sample.encode("utf-8")), ("same.tsv", reordered.getvalue().encode("utf-8"))],
                runtime_dir=Path(folder),
            )
        for filename in ("same.csv", "same.tsv"):
            finding = next(item for item in result["findings"]
                           if item["file"] == filename and item["kind"] == "identical_after_rounding"
                           and item["raw"]["rounded_to"] == 1.1)
            actual = {(cell["row_idx"], cell["column"], round(float(cell["original"]), 2))
                      for cell in finding["evidence"]["focus_cells"]}
            expected = {(2, "independent", 1.12), (8, "treated", 1.12),
                        (13, "control", 1.08), (25, "independent", 1.09),
                        (32, "treated", 1.06)}
            if filename == "same.tsv":
                expected = {(38 - row, column, value) for row, column, value in expected}
            self.assertEqual(expected, actual, filename)

    def test_valid_and_corrupt_material_is_partial_and_report_explains_omission(self):
        with tempfile.TemporaryDirectory() as folder:
            runtime = Path(folder)
            result = app.scan_files(
                [("valid.csv", _table_bytes()), ("broken.xlsx", b"not an Excel workbook")],
                runtime_dir=runtime,
            )
            materials = {item["filename"]: item for item in result["materials"]}
            self.assertEqual("read", materials["valid.csv"]["status"])
            self.assertEqual("failed", materials["broken.xlsx"]["status"])
            self.assertTrue(materials["broken.xlsx"]["reason_cn"])
            self.assertEqual(0, materials["broken.xlsx"]["numeric_cells"])
            self.assertEqual("partial", result["coverage_summary"]["status"])
            report = (runtime / result["scan_id"] / "report.html").read_text(encoding="utf-8")
            self.assertIn("broken.xlsx", report)
            self.assertIn(materials["broken.xlsx"]["reason_cn"], report)
            self.assertIn(result["coverage_summary"]["message_cn"], report)

    def test_all_corrupt_materials_preserve_failed_coverage_after_restart(self):
        with tempfile.TemporaryDirectory() as folder:
            runtime = Path(folder)
            job = app.start_scan_batch(
                [("broken.xlsx", b"not a workbook"), ("broken.docx", b"not a document")],
                runtime_dir=runtime,
            )
            current = _wait_for_job(job, runtime)
            self.assertEqual("succeeded", current["status"], current.get("error"))
            self.assertEqual("failed", current["result"]["scan_status"])
            self.assertEqual("failed", current["result"]["coverage_summary"]["status"])
            self.assertTrue(current["result"]["coverage_summary"]["message_cn"])
            self.assertEqual({"failed"}, {item["status"] for item in current["result"]["materials"]})
            self.assertTrue(all(item["reason_cn"] for item in current["result"]["materials"]))
            self.assertEqual([], current["result"]["findings"])
            persisted = _fresh_process(runtime, job["job_id"])["job"]
            self.assertEqual("failed", persisted["result"]["coverage_summary"]["status"])
            report = (runtime / persisted["result"]["scan_id"] / "report.html").read_text(encoding="utf-8")
            self.assertIn(persisted["result"]["coverage_summary"]["message_cn"], report)
            self.assertIn("broken.xlsx", report)
            self.assertIn("broken.docx", report)
            self.assertNotIn("未发现需要核对的数据。", report)

    def test_workbook_keeps_good_sheet_when_another_sheet_exceeds_read_limit(self):
        from openpyxl import Workbook
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet

        workbook = Workbook()
        good = workbook.active
        good.title = "good"
        records = list(csv.reader(io.StringIO(_table_bytes().decode("utf-8"))))
        good.append(records[0])
        for label, measurement in records[1:]:
            good.append([label, float(measurement)])
        oversized = workbook.create_sheet("oversized")
        oversized.append(["sample", "measurement"])
        for index in range(60):
            oversized.append([f"S{index}", index + 0.123])
        payload = io.BytesIO()
        workbook.save(payload)
        workbook.close()

        iter_rows = ReadOnlyWorksheet.iter_rows
        source_sheets_read = []

        def checked_rows(worksheet, *args, **kwargs):
            source_sheets_read.append(worksheet.title)
            self.assertNotEqual("oversized", worksheet.title,
                                "证据读取不能再次加载引擎已按大小上限跳过的工作表")
            return iter_rows(worksheet, *args, **kwargs)

        with tempfile.TemporaryDirectory() as folder:
            with patch.dict(os.environ, {"PAPERCONAN_MAX_CELLS": "100"}), \
                    patch.object(ReadOnlyWorksheet, "iter_rows", checked_rows):
                result = app.scan_files(
                    [("book.xlsx", payload.getvalue()), ("companion.csv", _table_bytes())],
                    runtime_dir=Path(folder),
                )
            materials = {item["filename"]: item for item in result["materials"]}
            self.assertEqual("partial", materials["book.xlsx"]["status"])
            self.assertEqual(1, materials["book.xlsx"]["sheets"])
            self.assertEqual(16, materials["book.xlsx"]["numeric_cells"])
            self.assertEqual("read", materials["companion.csv"]["status"])
            self.assertEqual("partial", result["coverage_summary"]["status"])
            self.assertEqual(["good"], source_sheets_read)
            pairs = [pair for finding in result["findings"]
                     for pair in finding["evidence"].get("comparison", [])]
            self.assertTrue(pairs, "正常工作表应保留跨文件原始数值证据")
            source_cells = [cell for pair in pairs for cell in (pair["left"], pair["right"])
                            if cell["file"] == "book.xlsx"]
            self.assertTrue(source_cells)
            self.assertEqual({"good"}, {cell["sheet"] for cell in source_cells})
            self.assertIn((2, 1.123), {(cell["row_idx"], float(cell["value"])) for cell in source_cells})


class BatchInputBoundaryTests(unittest.TestCase):
    def assert_rejected_before_writing(self, uploads, message):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(app.DemoError, message):
                app.start_scan_batch(uploads, runtime_dir=Path(folder))
            self.assertEqual([], list(Path(folder).iterdir()))

    def test_duplicate_names_are_case_insensitive(self):
        self.assert_rejected_before_writing(
            [("Data.csv", b"value\n1\n"), ("data.CSV", b"value\n2\n")], "重名|同名")

    def test_path_traversal_and_windows_special_names_are_rejected(self):
        for filename in ("../outside.csv", "..\\outside.csv", "C:\\outside.csv", "NUL.csv", "x.csv:stream"):
            with self.subTest(filename=filename):
                self.assert_rejected_before_writing([(filename, b"value\n1\n")], "文件名")

    def test_empty_batch_and_file_count_limit(self):
        self.assert_rejected_before_writing([], "选择|材料")
        uploads = [(f"file{index}.csv", b"v\n1\n") for index in range(app.MAX_FILES + 1)]
        self.assert_rejected_before_writing(uploads, "最多|材料")

    def test_per_file_size_limit_without_large_allocation(self):
        with patch.object(app, "MAX_UPLOAD_BYTES", 8):
            self.assert_rejected_before_writing([("large.csv", b"123456789")], "单个|超过")

    def test_total_size_limit_without_large_allocation(self):
        with patch.object(app, "MAX_BATCH_BYTES", 10):
            self.assert_rejected_before_writing(
                [("a.csv", b"v\n1.2\n"), ("b.csv", b"v\n2.3\n")], "合计|总|超过")


class PersistentHistoryTests(unittest.TestCase):
    def test_finished_batch_is_available_after_process_restart(self):
        with tempfile.TemporaryDirectory() as folder:
            runtime = Path(folder)
            job = app.start_scan_batch(
                [("experiment.csv", _table_bytes()), ("supplement.csv", _table_bytes())],
                runtime_dir=runtime,
            )
            current = _wait_for_job(job, runtime)
            self.assertEqual("succeeded", current["status"], current.get("error"))
            output = _fresh_process(runtime, job["job_id"], restore=True)
            self.assertEqual([job["job_id"]], [item["job_id"] for item in output["jobs"]])
            restored = output["job"]
            self.assertEqual("succeeded", restored["status"])
            self.assertEqual(["experiment.csv", "supplement.csv"], restored["result"]["sources"])
            self.assertEqual(current["result"]["scan_id"], restored["result"]["scan_id"])
            self.assertTrue(restored["result"]["findings"])
            self.assertNotIn("uploads", restored)
            self.assertNotIn("payload", restored)

    def test_interrupted_jobs_fail_on_restart_without_fabricated_result(self):
        for state in ("queued", "running"):
            with self.subTest(state=state), tempfile.TemporaryDirectory() as folder:
                runtime = Path(folder)
                job_id = "a" * 12
                job_store.save(runtime, {"job_id": job_id, "status": state,
                                        "created_at": "2026-09-23T00:00:00+00:00",
                                        "filename": "unfinished.csv", "scan_id": None})
                restored = _fresh_process(runtime, job_id, restore=True)["job"]
                self.assertEqual("failed", restored["status"])
                self.assertRegex(restored["error"], "中断|关闭")
                self.assertTrue(restored["finished_at"])
                self.assertIsNone(restored.get("result"))

    def test_old_result_import_is_idempotent_across_restarts(self):
        with tempfile.TemporaryDirectory() as folder:
            runtime = Path(folder)
            legacy_id = "b" * 12
            legacy_dir = runtime / legacy_id
            legacy_dir.mkdir()
            result = {"scan_id": legacy_id, "source": "legacy.csv", "profile": "review",
                      "scan_status": "complete", "summary": {"files": 1, "findings": 0},
                      "findings": [], "artifacts": {"user_report": f"/artifact/{legacy_id}/report.html"}}
            (legacy_dir / "result.json").write_text(json.dumps(result), encoding="utf-8")
            (legacy_dir / "report.html").write_text(
                "<!doctype html><html><body>Legacy screening report</body></html>", encoding="utf-8")
            first = _fresh_process(runtime, legacy_id, restore=True)
            second = _fresh_process(runtime, legacy_id, restore=True)
            self.assertEqual(1, len(first["jobs"]))
            self.assertEqual(1, len(second["jobs"]))
            self.assertEqual(first["jobs"][0]["job_id"], second["jobs"][0]["job_id"])
            self.assertEqual(result, second["job"]["result"])

    def test_legacy_import_preserves_partial_and_failed_scan_status(self):
        for status in ("partial", "failed"):
            with self.subTest(status=status), tempfile.TemporaryDirectory() as folder:
                runtime = Path(folder)
                legacy_id = "c" * 12
                legacy_dir = runtime / legacy_id
                legacy_dir.mkdir()
                result = {"scan_id": legacy_id, "source": "legacy.xlsx", "profile": "review",
                          "scan_status": status, "summary": {"files": 1, "findings": 0},
                          "findings": []}
                if status == "partial":
                    result["coverage_summary"] = {"status": status, "message_cn": "部分工作表未读取"}
                (legacy_dir / "result.json").write_text(json.dumps(result), encoding="utf-8")
                (legacy_dir / "report.html").write_text(
                    f"<!doctype html><html><body>{status}</body></html>", encoding="utf-8")
                restored = _fresh_process(runtime, legacy_id, restore=True)
                self.assertEqual(1, len(restored["jobs"]))
                self.assertEqual(status, restored["jobs"][0]["scan_status"])
                self.assertEqual(status, restored["job"]["scan_status"])
                self.assertEqual(status, restored["job"]["result"]["scan_status"])

    def test_result_without_report_is_not_imported_as_completed_history(self):
        with tempfile.TemporaryDirectory() as folder:
            runtime = Path(folder)
            unfinished_id = "d" * 12
            unfinished_dir = runtime / unfinished_id
            unfinished_dir.mkdir()
            result = {"scan_id": unfinished_id, "source": "unfinished.csv", "profile": "review",
                      "scan_status": "complete", "summary": {"files": 1, "findings": 0},
                      "findings": []}
            (unfinished_dir / "result.json").write_text(json.dumps(result), encoding="utf-8")
            self.assertEqual([], _fresh_process(runtime, restore=True)["jobs"])
            self.assertTrue((unfinished_dir / "result.json").exists())

    def test_report_write_failure_cannot_reappear_as_success_after_restart(self):
        write_text = Path.write_text

        def fail_report_write(path, *args, **kwargs):
            if path.name == "report.html":
                raise OSError("Injected report write failure")
            return write_text(path, *args, **kwargs)

        with tempfile.TemporaryDirectory() as folder:
            runtime = Path(folder)
            with patch.object(Path, "write_text", fail_report_write):
                job = app.start_scan_batch([("experiment.csv", _table_bytes())], runtime_dir=runtime)
                current = _wait_for_job(job, runtime)
            self.assertEqual("failed", current["status"])
            self.assertTrue(current["error"])
            self.assertEqual(job["job_id"], current["scan_id"])
            run_dir = runtime / current["scan_id"]
            self.assertTrue((run_dir / "result.json").is_file(), "测试必须实际到达报告写入阶段")
            self.assertFalse((run_dir / "report.html").exists())
            restored = _fresh_process(runtime, job["job_id"], restore=True)
            self.assertEqual([job["job_id"]], [item["job_id"] for item in restored["jobs"]])
            self.assertEqual("failed", restored["jobs"][0]["status"])
            self.assertEqual("failed", restored["job"]["status"])
            self.assertIsNone(restored["job"].get("result"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
