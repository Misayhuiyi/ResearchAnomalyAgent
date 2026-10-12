import json
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from platform_ports import LOCAL_ACTOR,Actor
from review_store import ReviewStore,ReviewError
from reference_comparison import ReferenceComparisonService

class ReferenceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        self.scan='abcdef123456';folder=self.root/self.scan/'input';folder.mkdir(parents=True)
        (folder/'original.csv').write_text('id,value\n1,0.1\n2,3.1\n3,1.5\n',encoding='utf-8')
        ReviewStore(self.root).register(LOCAL_ACTOR,{'scan_id':self.scan,'source':'original.csv','findings':[]})
        self.service=ReferenceComparisonService(self.root)
        self.preview=self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','reference.csv',b'id,value\n3,1.4\n1,0.2\n2,3.6\n')

    def compare(self,**extra):
        args={'source_sheet':'original','reference_sheet':'reference','keys':[{'source':0,'reference':0}],
              'columns':[{'source':1,'reference':1}],'confirm_order':False};args.update(extra)
        return self.service.compare(LOCAL_ACTOR,self.scan,self.preview['reference_id'],**args)

    def test_key_match_retains_both_physical_positions_and_exact_values(self):
        result=self.compare();self.assertEqual(result['differences_total'],3)
        cell=result['differences'][0];self.assertEqual(cell['source']['row'],2);self.assertEqual(cell['reference']['row'],3)
        self.assertEqual(cell['source']['value'],'0.1');self.assertEqual(cell['reference']['value'],'0.2')
        self.assertIn('不能',result['boundary']);json.dumps(result,allow_nan=False)
        self.assertEqual(self.compare()['comparison_id'],result['comparison_id'])

    def test_order_matching_requires_explicit_confirmation(self):
        with self.assertRaises(ReviewError):self.compare(keys=[])
        self.assertEqual(self.compare(keys=[],confirm_order=True)['matched_rows'],3)

    def test_duplicate_or_blank_key_is_rejected(self):
        for payload in (b'id,value\n1,.2\n1,3.6\n3,1.4\n',b'id,value\n1,.2\n,3.6\n3,1.4\n'):
            preview=self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','duplicate.csv',payload)
            with self.assertRaises(ReviewError):self.service.compare(LOCAL_ACTOR,self.scan,preview['reference_id'],'original','duplicate',
                [{'source':0,'reference':0}],[{'source':1,'reference':1}],False)

    def test_numeric_display_precision_is_not_false_difference(self):
        preview=self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','equal.csv',b'id,value\n1,0.10\n2,3.100\n3,1.50\n')
        result=self.service.compare(LOCAL_ACTOR,self.scan,preview['reference_id'],'original','equal',
            [{'source':0,'reference':0}],[{'source':1,'reference':1}],False)
        self.assertEqual(result['differences_total'],0)

    def test_filename_extension_mapping_and_owner_boundaries(self):
        for filename in ('../x.csv','..\\x.csv','x.exe','CON.csv'):
            with self.assertRaises(ReviewError):self.service.upload(LOCAL_ACTOR,self.scan,'original.csv',filename,b'1,2')
        with self.assertRaises(ReviewError):self.compare(columns=[{'source':99,'reference':1}])
        with self.assertRaises(ReviewError):self.service.view(Actor('other','other'),self.scan,self.preview['reference_id'])
        with self.assertRaises(ReviewError):self.service.view(LOCAL_ACTOR,self.scan,'../escape')

    def test_reference_context_is_source_bound_and_contains_actual_differences(self):
        result=self.compare();context=self.service.report_context(LOCAL_ACTOR,self.scan)
        self.assertEqual(context['scan_id'],self.scan);self.assertEqual(context['items'][0]['differences'],result['differences'])
        store=ReviewStore(self.root)
        with store.connect() as db:self.assertEqual(context['source_hash'],store._scan(db,LOCAL_ACTOR,self.scan)['source_hash'])

    def test_reference_report_in_all_formats_and_immutable_old_snapshot(self):
        from docx import Document
        from pypdf import PdfReader
        from openpyxl import load_workbook
        import report_exports
        store=ReviewStore(self.root);before=store.freeze(LOCAL_ACTOR,self.scan,0,'before-0001')
        self.compare();context=self.service.report_context(LOCAL_ACTOR,self.scan)
        snapshot=store.freeze(LOCAL_ACTOR,self.scan,0,'after-00001',reference_review=context)
        report_exports.export_snapshot(store,LOCAL_ACTOR,snapshot);folder=self.root/self.scan/'reports'/snapshot['report_id']
        doc=Document(folder/'report.docx');word='\n'.join(p.text for p in doc.paragraphs)
        pdf='\n'.join(p.extract_text() for p in PdfReader(folder/'report.pdf').pages)
        book=load_workbook(folder/'details.xlsx',read_only=True);excel='\n'.join(book.sheetnames)+'\n'+'\n'.join(str(c) for s in book for row in s.values for c in row);book.close()
        html=(folder/'report.html').read_text(encoding='utf-8')
        for output in (word,pdf,excel,html):
            for text in ('版本数值差异','original.csv','reference.csv','0.1','0.2','不能据此判断哪一版正确'):
                self.assertIn(text,output)
        self.assertNotIn('reference_review',store.snapshot(LOCAL_ACTOR,before['report_id']))

    def test_new_reference_quota_and_corrupt_workbook_fail_clearly(self):
        with self.assertRaises(ReviewError):self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','corrupt.xlsx',b'not a workbook')
        for _ in range(19):self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','same.csv',b'id,value\n1,.2\n')
        with self.assertRaises(ReviewError):self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','over.csv',b'id,value\n1,.2\n')

    def test_changed_source_and_uncached_formula_are_not_compared(self):
        from openpyxl import Workbook
        import io
        from table_sources import read_tables
        path=self.root/self.scan/'input/original.csv'
        path.write_text('id,value\n1,.9\n',encoding='utf-8')
        with self.assertRaises(ReviewError):self.compare()
        book=Workbook();book.active.append(['id','value']);book.active.append([1,'=1+1']);stream=io.BytesIO();book.save(stream);book.close()
        with self.assertRaises(ReviewError):self.service.upload(LOCAL_ACTOR,self.scan,'original.csv','formulas.xlsx',stream.getvalue())

if __name__=='__main__':unittest.main()
