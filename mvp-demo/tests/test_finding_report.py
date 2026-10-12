import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from platform_ports import LOCAL_ACTOR
from review_store import ReviewStore,Conflict
from finding_explanations import FindingExplanationService
from test_finding_explanations import PRICE,VALID
from model_gateway import ENDPOINT
import report_exports
import time

class FindingReportTests(unittest.TestCase):
    def test_bound_ai_is_frozen_and_exported_in_all_four_formats(self):
        from docx import Document
        from pypdf import PdfReader
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);store=ReviewStore(root);scan='aabbcc112233'
            source={'scan_id':scan,'source':'original.csv','findings':[{'id':'FD-0001','title_cn':'固定差值','summary_cn':'35行差值0.3','n':35,'evidence':{}}]}
            store.register(LOCAL_ACTOR,source)
            transport=Mock(return_value={'model':'deepseek-flash','usage':{'prompt_tokens':350,'completion_tokens':100},
                'choices':[{'finish_reason':'stop','message':{'content':json.dumps(VALID,ensure_ascii=False)}}]})
            service=FindingExplanationService(root,lambda:{'endpoint':ENDPOINT,'price':PRICE,'secret':'SYNTHETIC-NOT-A-KEY','valid_until':time.time()+3600},transport)
            quote=service.quote(LOCAL_ACTOR,scan,'FD-0001');service.run(LOCAL_ACTOR,scan,'FD-0001',quote['quote']['quote_id'],'approve-0001',True)
            context=service.report_context(LOCAL_ACTOR,scan)
            snap=store.freeze(LOCAL_ACTOR,scan,0,'freeze-0001',finding_ai=context)
            report_exports.export_snapshot(store,LOCAL_ACTOR,snap)
            folder=root/scan/'reports'/snap['report_id']
            doc=Document(folder/'report.docx');word='\n'.join(p.text for p in doc.paragraphs)
            pdf='\n'.join(p.extract_text() for p in PdfReader(folder/'report.pdf').pages)
            book=load_workbook(folder/'details.xlsx',read_only=True);excel='\n'.join(str(c) for s in book for row in s.values for c in row);book.close()
            html=(folder/'report.html').read_text(encoding='utf-8')
            for content in (word,pdf,excel,html):
                self.assertIn(VALID['what'],content);self.assertIn('不能判断作者意图',content)
            self.assertEqual(transport.call_count,1,'Export must never call model')
            self.assertEqual(store.snapshot(LOCAL_ACTOR,snap['report_id'])['finding_ai'],context)
            with self.assertRaises(Conflict):
                store.freeze(LOCAL_ACTOR,scan,0,'freeze-0002',finding_ai={**context,'source_hash':'0'*64})

if __name__=='__main__':unittest.main()
