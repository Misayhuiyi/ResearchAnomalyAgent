"""Synthetic text tables through actual file parsers and numerical engine."""
import io
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app


class FormatTests(unittest.TestCase):
    @staticmethod
    def rows():
        return [['control','treated']]+[[str(i/10),str(i/10+.3)] for i in [2,7,14,19,27,35,44,56,67,82]]

    def test_real_text_pdf_word_and_excel_tables(self):
        from reportlab.platypus import SimpleDocTemplate,Table,TableStyle
        from reportlab.lib import colors
        from docx import Document
        from openpyxl import Workbook
        rows=self.rows();pdf=io.BytesIO();table=Table(rows,colWidths=[130,130])
        table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black)]))
        SimpleDocTemplate(pdf).build([table])
        word=io.BytesIO();document=Document();t=document.add_table(rows=len(rows),cols=2)
        for i,row in enumerate(rows):
            for j,value in enumerate(row):t.cell(i,j).text=value
        document.save(word)
        excel=io.BytesIO();book=Workbook();book.active.title='Study'
        for index,row in enumerate(rows):book.active.append(row if index==0 else [float(value) for value in row])
        book.save(excel);book.close()
        with tempfile.TemporaryDirectory() as directory:
            result=app.scan_files([('paper.pdf',pdf.getvalue()),('paper.docx',word.getvalue()),('paper.xlsx',excel.getvalue())],runtime_dir=Path(directory))
            materials={m['filename']:m for m in result['materials']}
            for name in ['paper.pdf','paper.docx','paper.xlsx']:
                with self.subTest(file=name):
                    self.assertEqual(materials[name]['status'],'read',materials[name])
                    self.assertGreater(materials[name]['numeric_cells'],0)
                    self.assertTrue(any(f.get('file')==name for f in result['findings']),name)
            self.assertTrue(any('p1_t1' in f.get('sheet','') for f in result['findings'] if f.get('file')=='paper.pdf'))


if __name__=='__main__':unittest.main()
