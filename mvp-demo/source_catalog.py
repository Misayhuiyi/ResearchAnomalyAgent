"""Bounded, source-addressed material reading. No OCR and no network access.

The server runs this reader in a disposable process with a wall-clock deadline.
Material roles are confirmed by the user, never guessed by a model.
"""
import csv
import hashlib
import json
import os
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_CHUNKS = 600
MAX_TEXT = 240000
MAX_ZIP = 50 * 1024 * 1024


def limit_reader_memory():
    """Hard child-process ceiling, including compressed PDF expansion."""
    limit = 512 * 1024 * 1024
    if os.name != 'nt':
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
        return None
    import ctypes
    from ctypes import wintypes
    size = ctypes.c_size_t
    class Basic(ctypes.Structure):
        _fields_ = [('process_time',ctypes.c_int64),('job_time',ctypes.c_int64),('flags',wintypes.DWORD),
                    ('min_working',size),('max_working',size),('active',wintypes.DWORD),
                    ('affinity',size),('priority',wintypes.DWORD),('scheduling',wintypes.DWORD)]
    class IO(ctypes.Structure):
        _fields_ = [(name,ctypes.c_uint64) for name in ['read','write','other','read_bytes','write_bytes','other_bytes']]
    class Extended(ctypes.Structure):
        _fields_ = [('basic',Basic),('io',IO),('process_memory',size),('job_memory',size),('peak_process',size),('peak_job',size)]
    kernel = ctypes.WinDLL('kernel32',use_last_error=True)
    kernel.CreateJobObjectW.restype=wintypes.HANDLE
    kernel.CreateJobObjectW.argtypes=[ctypes.c_void_p,wintypes.LPCWSTR]
    kernel.SetInformationJobObject.argtypes=[wintypes.HANDLE,ctypes.c_int,ctypes.c_void_p,wintypes.DWORD]
    kernel.AssignProcessToJobObject.argtypes=[wintypes.HANDLE,wintypes.HANDLE]
    kernel.GetCurrentProcess.restype=wintypes.HANDLE
    handle=kernel.CreateJobObjectW(None,None)
    info=Extended();info.basic.flags=0x100;info.process_memory=limit
    if not handle or not kernel.SetInformationJobObject(handle,9,ctypes.byref(info),ctypes.sizeof(info)) or not kernel.AssignProcessToJobObject(handle,kernel.GetCurrentProcess()):
        raise RuntimeError('无法建立安全的材料读取资源限制。')
    return handle  # Keep the handle open until process exit; no unrelated process is attached.


def zip_guard(path):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_ZIP:
            raise ValueError('解压后的材料超过读取容量。')
        if any(i.file_size > 20*1024*1024 or i.flag_bits & 1 for i in infos):
            raise ValueError('材料包含超大内容或加密内容。')


def pieces(path, role):
    suffix = path.suffix.lower()
    if suffix in {'.csv', '.tsv'}:
        with path.open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.reader(handle, delimiter='\t' if suffix == '.tsv' else ',')
            headers = next(reader, [])
            for row, cells in enumerate(reader, 2):
                if row > 50002 or len(cells) > 200: raise ValueError('数据表超过读取容量。')
                yield f'行 {row}', role, ' | '.join(f'{h}: {v}' for h, v in zip(headers, cells))
    elif suffix in {'.xlsx', '.xlsm'}:
        import openpyxl
        zip_guard(path)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                if sheet.max_row > 50000 or sheet.max_column > 200: raise ValueError('数据表超过读取容量。')
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, [])
                for row, cells in enumerate(rows, 2):
                    yield f'{sheet.title} · 行 {row}', role, ' | '.join(f'{h}: {v}' for h, v in zip(headers, cells) if v is not None)
        finally: book.close()
    elif suffix == '.docx':
        zip_guard(path)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with zipfile.ZipFile(path) as archive:
            tree = ET.fromstring(archive.read('word/document.xml'))
        body = tree.find('w:body', ns)
        for index, element in enumerate(body, 1):
            if element.tag.endswith('}tbl'):
                rows = element.findall('w:tr', ns)
                header = ''
                for row, node in enumerate(rows, 1):
                    cells = [''.join(t.text or '' for t in cell.findall('.//w:t', ns)) for cell in node.findall('w:tc', ns)]
                    if row == 1: header = ' | '.join(cells)
                    yield f'表格块 {index} · 行 {row}', 'table' if role == 'body' else role, header+'\n'+' | '.join(cells)
            else:
                text = ''.join(t.text or '' for t in element.findall('.//w:t', ns))
                yield f'正文块 {index}', role, text
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted or len(reader.pages) > 200: raise ValueError('PDF 已加密或超过 200 页。')
        for index, page in enumerate(reader.pages, 1):
            # Text only: page content is not promoted to a table without a human mapping.
            text = page.extract_text() or ''
            yield f'第 {index} 页', role, text
    else:
        raise ValueError('此格式尚不能用于论文内容核对，请提供 CSV、Excel、Word 或文字 PDF。')


def build(root, manifest, roles):
    catalog, coverage = {}, []
    size = 0
    for item in manifest:
        name = item['filename']
        if Path(name).name != name or roles.get(name) not in {'body', 'table', 'supplement'}:
            raise ValueError('请确认每份材料的用途。')
        path = root / name
        if hashlib.sha256(path.read_bytes()).hexdigest() != item['sha256']:
            raise ValueError('原始材料校验未通过，不能继续。')
        count, status, reason = 0, 'read', ''
        try:
            for location, kind, text in pieces(path, roles[name]):
                if not text.strip(): continue
                # Exact substrings preserve offsets; splitting is declared in each location.
                for start in range(0, len(text), 1600):
                    value = text[start:start+1600]
                    if len(catalog) >= MAX_CHUNKS or size + len(value) > MAX_TEXT:
                        raise ValueError('本次内容读取达到容量上限，未读取的部分不参与核对。')
                    source_id = 's'+hashlib.sha256((name+'\0'+location+'\0'+str(start)).encode()).hexdigest()[:20]
                    catalog[source_id] = {'file': name, 'location': location + (f' · 字符 {start+1}–{start+len(value)}' if len(text)>1600 else ''),
                                          'kind': kind, 'text': value}
                    count += 1; size += len(value)
            if not count: status, reason = 'unread', '未读取到文字；扫描件与图片不在本版范围内。'
        except Exception:
            status = 'partial' if count else 'unread'
            reason = '材料格式、编码或读取容量不满足要求；仅已列出的片段可参与核对。'
        coverage.append({'file': name, 'status': status, 'source_count': count, 'reason': reason})
    return {'catalog': catalog, 'coverage': coverage}


if __name__ == '__main__':
    memory_guard = limit_reader_memory()
    request = json.loads(sys.stdin.read(65536))
    answer = build(Path(request['root']), request['manifest'], request['roles'])
    sys.stdout.write(json.dumps(answer, ensure_ascii=False))
