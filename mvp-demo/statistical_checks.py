"""Opt-in arithmetic checks with explicit context and conservative rounding bounds.

DET-05a: count / denominator * 100. DET-06a: SE = SD / sqrt(n).
The latter applies only to an ordinary independent-sample mean, not adjusted SEs.
Reference: https://www.itl.nist.gov/div898/handbook/eda/section3/eda352.htm
"""
from __future__ import annotations

import csv
import re
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, ROUND_CEILING, ROUND_FLOOR, localcontext
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from scipy import stats
import pattern_checks


MAX_ROWS = 50000
MAX_COLUMNS = 200
MAX_CELLS = 1000000
MAX_FINDINGS = 500
RULES = {
    "count_percent": {"id": "DET-05a", "name": "人数与百分比", "roles": ("count", "total", "percent"),
                      "assumption": "每行人数、总人数和百分比属于同一人群与时点，百分比按未加权人数计算。"},
    "sd_se": {"id": "DET-06a", "name": "标准差与标准误", "roles": ("n", "sd", "se"),
              "assumption": "每行样本量、标准差和均值标准误属于同一组独立观测，未使用加权、聚类或其他修正。"},
    "group_total": {"id": "DET-05b", "name": "分组合计", "roles": ("total",),
                    "assumption": "各分组属于同一人群、同一时点，互不重叠且包含全部人员。"},
    "participant_flow": {"id": "DET-05c", "name": "人数流转", "roles": ("initial", "excluded", "missing", "analysed"),
                         "assumption": "排除、缺失与最终分析人员互不重叠，三者合计为同一批初始人员。"},
    "mean_ci": {"id": "DET-06b", "name": "均值置信区间", "roles": ("n", "mean", "sd", "ci_low", "ci_high", "confidence", "ci_method"),
                "assumption": "区间是同组独立观测均值的双侧正态近似或t区间，材料已写明方法与置信水平，未经加权或校正。"},
    "reported_p": {"id": "DET-06c", "name": "统计量与P值", "roles": ("test", "statistic", "p", "tail"),
                   "assumption": "统计量与P值来自同一次未校正的z、t、卡方或F检验，材料已写明方向及所需自由度。"},
    "integer_mean": {"id": "DET-06d", "name": "整数评分均值", "roles": ("n", "mean"),
                     "assumption": "每条原始观测都是整数评分或计数，均值未经加权、转换或跨题平均，样本量为实际观测数。"},
}
ALIASES = {
    "count": {"人数", "事件人数", "事件数", "count", "event_count", "events"},
    "total": {"总人数", "总数", "total", "total_n", "n_total", "denominator"},
    "percent": {"百分比", "百分比(%)", "占比(%)", "percent", "percentage", "percent(%)", "percentage(%)"},
    "n": {"样本量", "样本数", "n", "sample_size"},
    "sd": {"标准差", "sd", "standard_deviation"},
    "se": {"标准误", "标准误差", "均值标准误", "se", "sem", "standard_error"},
    "initial": {"初始人数", "initial"}, "excluded": {"排除人数", "excluded"},
    "missing": {"缺失人数", "missing"}, "analysed": {"分析人数", "analysed", "analyzed"},
    "mean": {"均值", "mean"}, "ci_low": {"区间下限", "ci_low"}, "ci_high": {"区间上限", "ci_high"},
    "confidence": {"置信水平(%)", "置信水平", "confidence"}, "ci_method": {"区间方法", "ci_method"},
    "test": {"检验方法", "test"}, "statistic": {"统计量", "statistic"}, "p": {"p", "p值", "p_value"},
    "tail": {"检验方向", "tail"}, "df": {"自由度", "df", "分子自由度"}, "df2": {"分母自由度", "df2"},
}
TEXT_ROLES = {"ci_method", "test", "tail", "p"}
EXACT_ROLES = {"count", "total", "n", "initial", "excluded", "missing", "analysed", "df", "df2"}
for _mode, _definition in pattern_checks.MODES.items():
    RULES['patterns_'+_mode] = {'id':_definition['ids'][0], 'name':_definition['name'],
                              'roles':(), 'assumption':_definition['assumption']}


class CannotCheck(ValueError):
    pass


@dataclass(frozen=True)
class Cell:
    original: str
    value: Decimal
    quantum: Decimal | None


def selections(value: Any) -> list[str]:
    if value is None:
        return []
    if not isinstance(value, list) or any(not isinstance(item, str) or item not in RULES for item in value):
        raise CannotCheck("统计核对选项无效，请重新选择。")
    if len(value) != len(set(value)):
        raise CannotCheck("统计核对选项不能重复。")
    return value


def number(value: Any, fmt: str | None = None, percent: bool = False) -> Cell:
    if value is None or isinstance(value, bool):
        raise CannotCheck("存在空值或非数值")
    text = str(value).strip()
    if not text or len(text) > 80:
        raise CannotCheck("存在空值或非数值")
    has_percent = text.endswith(("%", "％"))
    if has_percent and not percent:
        raise CannotCheck("数值单位不适用")
    raw = text[:-1].strip() if has_percent else text
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d{1,2})?", raw):
        raise CannotCheck("存在空值、区间或非数值")
    try:
        parsed = Decimal(raw)
    except InvalidOperation as exc:
        raise CannotCheck("存在无效数值") from exc
    if not parsed.is_finite() or abs(parsed) > Decimal("1e12") or abs(parsed.as_tuple().exponent) > 12:
        raise CannotCheck("数值或精度超出本次核对范围")
    quantum = Decimal(1).scaleb(parsed.as_tuple().exponent)
    # CSV and textual spreadsheet cells retain the literal reporting precision.
    if fmt is None or isinstance(value, str):
        return Cell(text, parsed, quantum)
    match = re.fullmatch(r"(?:#,##)?0(?:\.(0{1,12}))?(%)?", fmt)
    if not match:
        return Cell(text, parsed, None)
    if match.group(2) and not percent:
        raise CannotCheck("数值单位不适用")
    display_quantum = Decimal(1).scaleb(-len(match.group(1) or ""))
    display_value = parsed * 100 if match.group(2) else parsed
    displayed = display_value.quantize(display_quantum, rounding=ROUND_HALF_UP)
    label = f"{text}（显示 {displayed}{'%' if match.group(2) else ''}）"
    return Cell(label, displayed, display_quantum)


def integer(cell: Cell, minimum: int = 0) -> int:
    if cell.value != cell.value.to_integral_value() or not minimum <= cell.value <= 1000000000:
        raise CannotCheck("人数或样本量不满足本项核对要求")
    return int(cell.value)


def interval(cell: Cell, minimum: Decimal | None = Decimal(0)) -> tuple[Decimal, Decimal]:
    if cell.quantum is None:
        raise CannotCheck("未明确报告的小数位数")
    lower = cell.value - cell.quantum / 2
    return (max(minimum, lower) if minimum is not None else lower), cell.value + cell.quantum / 2


def printable(value: Decimal) -> str:
    # This is a labelled calculated approximation; original cell text is kept separately.
    return format(value, ".10g")


def calculate(key: str, cells: dict[str, Cell]) -> dict:
    with localcontext() as context:
        context.prec = 50
        if key not in {"count_percent", "sd_se"}:
            return calculate_extended(key, cells)
        if key == "count_percent":
            count, total = integer(cells["count"]), integer(cells["total"], 1)
            if count > total or not 0 <= cells["percent"].value <= 100:
                raise CannotCheck("人数超出总人数或百分比超出0至100")
            expected = Decimal(count) / Decimal(total) * 100
            lower, upper = interval(cells["percent"])
            mismatch = not lower <= expected <= upper
            summary = f"{count} ÷ {total} × 100% ≈ {printable(expected)}%，材料中的百分比为 {cells['percent'].original}；考虑保留位数后仍不一致。"
            return {"mismatch": mismatch, "formula": "人数 ÷ 总人数 × 100%", "expected": str(expected),
                    "reported_interval": [str(lower), str(upper)], "summary": summary}
        n = integer(cells["n"], 2)
        if cells["sd"].value < 0 or cells["se"].value < 0:
            raise CannotCheck("标准差或标准误不能为负数")
        root = Decimal(n).sqrt()
        sd_low, sd_high = interval(cells["sd"])
        se_low, se_high = interval(cells["se"])
        expected = cells["sd"].value / root
        low, high = sd_low / root, sd_high / root
        mismatch = high < se_low or low > se_high
        summary = f"标准差 {cells['sd'].original} ÷ √{n}，推算标准误约 {printable(expected)}；材料中的标准误为 {cells['se'].original}，考虑两者保留位数后仍不一致。"
        return {"mismatch": mismatch, "formula": "均值标准误 = 标准差 ÷ √样本量", "expected": str(expected),
                "expected_interval": [str(low), str(high)], "reported_interval": [str(se_low), str(se_high)], "summary": summary}


def disjoint(left, right) -> bool:
    # SciPy CDF/quantiles are floating point. Guard numerical error, not reporting precision.
    guard = Decimal('1e-12')
    return left[1] + guard < right[0] or right[1] + guard < left[0]


def p_interval(text: str | Cell) -> tuple[Decimal, Decimal]:
    match = re.fullmatch(r'\s*([<>]=?|≤|≥)?\s*(.*?)\s*', text) if isinstance(text, str) else None
    cell = number(match.group(2)) if match else text
    if not 0 <= cell.value <= 1:
        raise CannotCheck('P值须在0至1之间')
    op = match.group(1) if match else None
    if op in {'<', '<=', '≤'}:
        return Decimal(0), cell.value
    if op in {'>', '>=', '≥'}:
        return cell.value, Decimal(1)
    lo, hi = interval(cell)
    return lo, min(Decimal(1), hi)


def calculate_extended(key: str, cells: dict) -> dict:
    if key in {'group_total', 'participant_flow'}:
        if key == 'group_total':
            parts = [integer(cell) for role, cell in cells.items() if role.startswith('group_')]
            total = integer(cells['total'])
            formula = '互斥完整分组人数合计 = 总人数'
        else:
            parts = [integer(cells[role]) for role in ['excluded', 'missing', 'analysed']]
            total = integer(cells['initial'])
            formula = '排除人数 + 缺失人数 + 分析人数 = 初始人数'
        return {'mismatch': sum(parts) != total, 'formula': formula, 'expected': str(sum(parts)),
                'summary': f"{' + '.join(map(str, parts))} = {sum(parts)}，材料中的总人数为 {total}，两者不一致。"}
    if key == 'integer_mean':
        n = integer(cells['n'], 1)
        lo, hi = interval(cells['mean'], None)
        total_lo, total_hi = lo * n, hi * n
        mismatch = total_lo.to_integral_value(rounding=ROUND_CEILING) > total_hi.to_integral_value(rounding=ROUND_FLOOR)
        return {'mismatch': mismatch, 'formula': '整数观测之和必须为整数，均值 × 样本量须与之相容',
                'expected_interval': [str(total_lo), str(total_hi)],
                'summary': f"样本量 {n}、均值 {cells['mean'].original}；按报告小数位反推的总分区间不包含整数，请核对原始评分。"}
    if key == 'mean_ci':
        n = integer(cells['n'], 2)
        if cells['sd'].value < 0 or not 0 < cells['confidence'].value < 100:
            raise CannotCheck('标准差或置信水平不适用；置信水平按百分数填写')
        method = cells['ci_method'].strip().lower()
        confidence = float(cells['confidence'].value / 100)
        if method in {'normal', 'z', '正态近似'}:
            critical = stats.norm.ppf((1 + confidence) / 2)
        elif method in {'t', 'student-t', 't方法'}:
            critical = stats.t.ppf((1 + confidence) / 2, n - 1)
        else:
            raise CannotCheck('区间方法未说明或不属于已登记的正态近似、t方法')
        multiplier = Decimal(str(critical)) / Decimal(n).sqrt()
        mean_low, mean_high = interval(cells['mean'], None)
        sd_low, sd_high = interval(cells['sd'])
        low_range = (mean_low - multiplier * sd_high, mean_high - multiplier * sd_low)
        high_range = (mean_low + multiplier * sd_low, mean_high + multiplier * sd_high)
        if cells['ci_low'].value > cells['ci_high'].value:
            raise CannotCheck('区间下限大于上限')
        mismatch = disjoint(low_range, interval(cells['ci_low'], None)) or disjoint(high_range, interval(cells['ci_high'], None))
        center, margin = cells['mean'].value, multiplier * cells['sd'].value
        return {'mismatch': mismatch, 'formula': '均值 ± 临界值 × 标准差 ÷ √样本量',
                'expected_interval': [str(center - margin), str(center + margin)],
                'lower_endpoint_range': list(map(str, low_range)), 'upper_endpoint_range': list(map(str, high_range)),
                'summary': f"按 {method} 方法和 {cells['confidence'].original}% 置信水平推算，区间约为 [{printable(center-margin)}, {printable(center+margin)}]；与材料区间考虑舍入后仍不一致。"}
    if key != 'reported_p':
        raise CannotCheck('该检查方法尚未登记')
    method = cells['test'].strip().lower()
    tail = {'双侧':'two-sided', '左侧':'left', '右侧':'right'}.get(cells['tail'].strip(), cells['tail'].strip().lower())
    if tail not in {'two-sided', 'left', 'right'}:
        raise CannotCheck('检验方向未明确，不能从P值推测单双侧')
    def degrees(role):
        if role not in cells or not 0 < cells[role].value <= 10000000:
            raise CannotCheck('缺少适用的自由度')
        return float(cells[role].value)
    if method == 'z':
        distribution = stats.norm
    elif method == 't':
        distribution = stats.t(degrees('df'))
    elif method in {'chi2', '卡方'}:
        distribution = stats.chi2(degrees('df'))
    elif method == 'f':
        distribution = stats.f(degrees('df'), degrees('df2'))
    else:
        raise CannotCheck('检验方法不属于已登记的z、t、卡方或F检验')
    lo, hi = interval(cells['statistic'], None)
    if method in {'chi2', '卡方', 'f'}:
        if cells['statistic'].value < 0 or tail != 'right':
            raise CannotCheck('卡方与F核对仅支持非负统计量的右尾检验')
        lo = max(lo, Decimal(0))
    if tail == 'two-sided':
        abs_lo = Decimal(0) if lo <= 0 <= hi else min(abs(lo), abs(hi))
        abs_hi = max(abs(lo), abs(hi))
        expected = (2 * distribution.sf(float(abs_hi)), 2 * distribution.sf(float(abs_lo)))
    elif tail == 'left':
        expected = (distribution.cdf(float(lo)), distribution.cdf(float(hi)))
    else:
        expected = (distribution.sf(float(hi)), distribution.sf(float(lo)))
    expected = tuple(Decimal(str(value)) for value in expected)
    if any(not value.is_finite() for value in expected):
        raise CannotCheck('统计分布计算未取得有效结果')
    reported = p_interval(cells['p'])
    return {'mismatch': disjoint(expected, reported), 'formula': f'{method} 分布，{tail}，按统计量舍入区间复算P值',
            'expected_interval': list(map(str, expected)), 'reported_interval': list(map(str, reported)),
            'summary': f"按材料中的统计量与检验方法，P值约在 {printable(expected[0])} 至 {printable(expected[1])}；与报告的 {cells['p'].original if isinstance(cells['p'], Cell) else cells['p']} 不相容。"}


def _header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")").lower()


def mapping(headers: list, key: str) -> dict[str, int]:
    result = {}
    for role in RULES[key]["roles"]:
        indices = [index for index, value in enumerate(headers) if _header(value) in ALIASES[role]]
        if len(indices) != 1:
            raise CannotCheck("缺少明确的列名，或同一含义有多列，无法自动对应")
        result[role] = indices[0]
    if key == 'group_total':
        for index, value in enumerate(headers):
            match = re.fullmatch(r'(?:分组([1-9]|1\d|20)人数|group_([1-9]|1\d|20))', _header(value))
            if match:
                role = 'group_' + (match.group(1) or match.group(2))
                if role in result:
                    raise CannotCheck('分组列重复，无法确认合计口径')
                result[role] = index
        if len(result) < 3:
            raise CannotCheck('至少需要两列明确的分组人数')
    if key == 'reported_p':
        for role in ('df', 'df2'):
            indices = [index for index, value in enumerate(headers) if _header(value) in ALIASES[role]]
            if len(indices) > 1:
                raise CannotCheck('自由度列存在歧义')
            if indices:
                result[role] = indices[0]
    return result


def table_check(filename: str, sheet: str, headers: list, rows: list, keys: list[str], limited: bool = False) -> tuple[list, list]:
    findings, coverage = [], []
    for key in keys:
        if key.startswith('patterns_'):
            found, covered = pattern_checks.check_table(filename, sheet, headers, rows, key.removeprefix('patterns_'), number)
            if limited:
                for item in covered:
                    if item['status']=='checked':
                        item['status']='partial'
                    item['notes'].append('源表超出读取容量，仅检查已读取范围。')
            findings.extend(found);coverage.extend(covered)
            continue
        rule = RULES[key]
        info = {"rule_id": rule["id"], "name": rule["name"], "file": filename, "sheet": sheet,
                "status": "insufficient", "checked_rows": 0, "skipped_rows": 0,
                "findings_total": 0, "findings_shown": 0, "assumption": rule["assumption"], "notes": []}
        coverage.append(info)
        try:
            columns = mapping(headers, key)
        except CannotCheck as exc:
            info["notes"].append(str(exc))
            if limited:
                info["notes"].append("表格超出统计核对容量，未完整读取。")
            info["skipped_rows"] = len(rows)
            continue
        reasons = Counter()
        for row in rows:
            values, formats = row["values"], row.get("formats")
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            try:
                if len(values) != len(headers):
                    raise CannotCheck("行列数量与表头不一致")
                cells = {}
                for role, index in columns.items():
                    if role in {'df', 'df2'} and (values[index] is None or str(values[index]).strip() == ''):
                        continue  # z needs neither; t/chi2 need df; F needs both.
                    fmt = formats[index] if formats and role not in EXACT_ROLES and not role.startswith('group_') else None
                    if role == 'p' and not isinstance(values[index], str):
                        cells[role] = number(values[index], fmt)
                    else:
                        cells[role] = ('' if values[index] is None else str(values[index]).strip()) if role in TEXT_ROLES else number(values[index], fmt, role in {'percent', 'confidence'})
                answer = calculate(key, cells)
            except (CannotCheck, InvalidOperation) as exc:
                info["skipped_rows"] += 1
                reasons[str(exc) or "无法计算"] += 1
                continue
            info["checked_rows"] += 1
            if not answer["mismatch"]:
                continue
            info["findings_total"] += 1
            if info["findings_shown"] >= MAX_FINDINGS:
                continue
            info["findings_shown"] += 1
            original_headers = [str(headers[index]) for index in columns.values()]
            findings.append({"id": "", "kind": key + "_inconsistent", "rule_id": rule["id"], "rule_version": "1.0",
                             "category": "statistical_consistency", "severity": "medium", "file": filename,
                             "sheet": sheet, "columns": original_headers, "n": 1,
                             "location": f"{filename} · {sheet} · 第 {row['row_idx']} 行",
                             "title_cn": rule["name"] + "不一致", "summary_cn": answer["summary"],
                             "rule": answer["formula"], "calculation": answer,
                             "next_steps": ["核对所列原始值及表格保留位数。", rule["assumption"]],
                             "evidence": {"headers": original_headers, "highlight_cols": list(range(1, len(columns) + 1)),
                                          "rows": [{"row_idx": row["row_idx"], "values": [cells[role].original if isinstance(cells.get(role), Cell) else str(cells.get(role, '')) for role in columns]}],
                                          "note_cn": "仅按已确认的计算口径提示待核对项，不修改原始材料。"}})
        info["notes"] = [f"{reason}：{count} 行" for reason, count in reasons.items()]
        if limited:
            info["notes"].append("仅核对本次读取范围，表格超出统计核对容量。")
        if info["findings_total"] > info["findings_shown"]:
            info["notes"].append(f"共命中 {info['findings_total']} 行，展示前 {info['findings_shown']} 行。")
        if info["checked_rows"]:
            info["status"] = "partial" if info["skipped_rows"] or limited else "checked"
        elif not info["notes"]:
            info["notes"].append("没有可核对的数据行。")
    return findings, coverage


def _take_rows(iterator, excel: bool, cell_limit: int | None = None) -> tuple[list, list, bool]:
    first = next(iterator, None)
    if first is None:
        return [], [], False
    headers = [cell.value for cell in first] if excel else list(first)
    if len(headers) > MAX_COLUMNS:
        return headers[:MAX_COLUMNS], [], True
    rows, cells = [], len(headers)
    capacity = MAX_CELLS if cell_limit is None else min(MAX_CELLS, cell_limit)
    for number_, row in enumerate(iterator, start=2):
        if len(rows) >= MAX_ROWS or cells + len(row) > capacity:
            return headers, rows, True
        if len(row) > MAX_COLUMNS:
            return headers, rows, True
        cells += len(row)
        rows.append({"row_idx": number_, "values": [cell.value for cell in row] if excel else list(row),
                     "formats": [cell.number_format for cell in row] if excel else None})
    return headers, rows, False


def check_file(path: Path, allowed_sheets: set[str], keys: list[str]) -> tuple[list, list]:
    findings, coverage = [], []
    if not keys:
        return findings, coverage

    def skipped(note: str, status: str = "insufficient", sheet: str = ""):
        return [], [{"rule_id": rule_id, "name": pattern_checks.NAMES.get(rule_id, RULES[key]["name"]), "file": path.name, "sheet": sheet,
                     "status": status, "checked_rows": 0, "skipped_rows": 0, "findings_total": 0,
                     "findings_shown": 0, "assumption": RULES[key]["assumption"], "notes": [note]} for key in keys
                    for rule_id in (pattern_checks.MODES[key[9:]]['ids'] if key.startswith('patterns_') else [RULES[key]['id']])]
    if not allowed_sheets:
        return skipped("材料未被完整读取，未进行统计核对。")
    try:
        if path.suffix.lower() in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as handle:
                headers, rows, limited = _take_rows(iter(csv.reader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")), False)
            return table_check(path.name, path.stem, headers, rows, keys, limited)
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            workbook = load_workbook(path, data_only=True, read_only=True)
            try:
                for sheet in workbook.worksheets:
                    if sheet.title not in allowed_sheets:
                        continue
                    if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                        coverage.extend(skipped("超出统计核对容量，未重新读取。", sheet=sheet.title)[1])
                        continue
                    headers, rows, limited = _take_rows(iter(sheet.iter_rows()), True)
                    found, covered = table_check(path.name, sheet.title, headers, rows, keys, limited)
                    findings.extend(found)
                    coverage.extend(covered)
            finally:
                workbook.close()
            return (findings, coverage) if coverage else skipped("没有可对应的已读取数据表。")
        return skipped("当前统计核对支持 CSV、TSV、XLSX、XLSM，其他格式尚未核对。", "unsupported")
    except (OSError, UnicodeError, csv.Error, ValueError, BadZipFile, KeyError) as exc:
        return skipped("无法读取统计核对所需的原始值或显示格式。", "failed")
