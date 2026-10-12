"""Bounded table reader retaining physical row/column coordinates and raw values."""
import csv
import math
from pathlib import Path
from source_catalog import zip_guard

MAX_CELLS=200000
MAX_ROWS=50000
MAX_COLUMNS=200
FORMATS={'.csv','.tsv','.xlsx','.xlsm'}

def number(value):
    if value is None or isinstance(value,bool):return False
    try:return math.isfinite(float(value))
    except (ValueError,TypeError):return False

def recover_headers(table):
    """Only skip sparse title rows, never infer labels from measurements."""
    first=list(table.get('headers') or [])
    candidates=[(1,first)]+[(r['row_idx'],r['values']) for r in table.get('rows',[])[:5]]
    def labels(values):return sum(v is not None and str(v).strip()!='' and not number(v) for v in values)
    valid=[(idx,list(values)) for idx,values in candidates if labels(values)>=2 and not any(number(v) for v in values)]
    if labels(first)<=1 and valid:
        index,headers=max(valid,key=lambda p:labels(p[1]))
    else:index,headers=1,first
    table['headers']=[str(v).strip() if v is not None and str(v).strip() else f'第 {i+1} 列' for i,v in enumerate(headers)]
    table['header_row']=index
    return table

def take(iterator):
    first=next(iterator,None)
    if first is None:return {'headers':[],'rows':[],'header_row':1,'limited':False}
    headers=list(first)
    if len(headers)>MAX_COLUMNS:raise ValueError('参考表超过 200 列，请拆分。')
    if any(isinstance(v,str) and len(v)>4096 for v in headers):raise ValueError('参考表含超长内容，请缩小范围。')
    rows=[];count=len(headers)
    for index,row in enumerate(iterator,2):
        values=list(row);count+=len(values)
        if any(isinstance(v,float) and not math.isfinite(v) for v in values):raise ValueError('参考表含无效数值，请修正后核对。')
        if any(isinstance(v,str) and len(v)>4096 for v in values):raise ValueError('参考表含超长内容，请缩小范围。')
        if index>MAX_ROWS+1 or len(values)>MAX_COLUMNS or count>MAX_CELLS:
            raise ValueError('版本核对每张表最多 50000 行、20 万单元格，请拆分后核对。')
        rows.append({'row_idx':index,'values':values})
    return recover_headers({'headers':headers,'rows':rows,'limited':False})

def read_tables(path):
    path=Path(path)
    if path.suffix.lower() not in FORMATS:raise ValueError('版本核对支持 CSV、TSV、XLSX 和 XLSM；其他格式请先另存为这些格式。')
    if path.suffix.lower() in {'.csv','.tsv'}:
        with path.open(encoding='utf-8-sig',newline='') as handle:
            return {path.stem:take(iter(csv.reader(handle,delimiter='\t' if path.suffix.lower()=='.tsv' else ',')))}
    from openpyxl import load_workbook
    zip_guard(path);book=load_workbook(path,read_only=True,data_only=True)
    formulas=None
    try:
        formulas=load_workbook(path,read_only=True,data_only=False)
        if len(book.worksheets)>20:raise ValueError('版本核对最多 20 张工作表，请拆分文件。')
        tables={};total=0
        for sheet in book:
            if (sheet.max_column or 0)>MAX_COLUMNS:raise ValueError('参考表超过 200 列，请拆分。')
            table=take(iter(sheet.values));total+=len(table['headers'])+sum(len(r['values']) for r in table['rows'])
            for formula_row,cached_row in zip(formulas[sheet.title].values,sheet.values):
                for col,formula in enumerate(formula_row):
                    if isinstance(formula,str) and formula.startswith('=') and (col>=len(cached_row) or cached_row[col] is None):
                        raise ValueError('表格含未保存计算结果的公式，请在 Excel 重算并保存后再核对。')
            if total>MAX_CELLS:raise ValueError('版本核对每份文件最多 20 万单元格，请拆分。')
            tables[sheet.title]=table
        return tables
    finally:
        book.close()
        if formulas:formulas.close()
