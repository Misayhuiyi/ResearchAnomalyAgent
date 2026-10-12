from __future__ import annotations

import argparse
import base64
import binascii
import csv
import html
import hashlib
import ipaddress
import json
import mimetypes
import math
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse
from zipfile import BadZipFile

import job_store
import public_materials
import statistical_checks
import review_store
import report_exports
import agent_service
import finding_explanations
import reference_comparison
from table_sources import recover_headers
from budget_store import BudgetError
from model_gateway import ModelError
from evidence import normalise_cross_finding
from platform_ports import LOCAL_ACTOR, LocalMaterialStore
from public_fetch import MaterialError
from runtime_guard import RuntimeLease


ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
SAMPLES_DIR = ROOT / "samples"
RUNTIME_DIR = Path(os.environ.get("PAPER_AUDIT_RUNTIME_DIR", str(ROOT / "runtime" / "scans")))
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_BATCH_BYTES = 100 * 1024 * 1024
MAX_FILES = 20
ALLOWED_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls", ".xlsm", ".xlsb", ".pdf", ".docx"}
ALLOWED_PROFILES = {"review", "triage", "forensic"}
SAFE_SCAN_ID = re.compile(r"^[a-f0-9]{12}$")
JOB_EXECUTOR = ThreadPoolExecutor(max_workers=2, thread_name_prefix="paper-audit-scan")
JOB_LOCK = threading.Lock()
SUBMISSION_LOCK = threading.Lock()
JOBS: dict[str, dict[str, Any]] = {}
JOB_CANCEL: dict[str, threading.Event] = {}


class DemoError(Exception):
    """Expected request or scan error safe to display to the user."""


class ScanCancelled(DemoError):
    pass


ENGINE_RSS_LIMIT = 1024 * 1024 * 1024


def run_engine(command, cancel_event=None):
    import psutil
    cancel_event = cancel_event or threading.Event()
    if cancel_event.is_set():
        raise ScanCancelled('检查已停止，原始材料仍保留。')
    started = time.monotonic()
    with subprocess.Popen(command, cwd=ROOT, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                          text=True, encoding='utf-8', errors='replace') as process:
        owner=psutil.Process(process.pid)
        known={owner.pid:owner}
        def stop_tree():
            try: known.update({p.pid:p for p in owner.children(recursive=True)})
            except psutil.NoSuchProcess: pass
            for child in reversed(list(known.values())):
                try:
                    if child.is_running(): child.kill()
                except psutil.NoSuchProcess: pass
            process.communicate(timeout=10)
        while True:
            if cancel_event.is_set() or time.monotonic()-started > 180:
                stop_tree()
                if cancel_event.is_set(): raise ScanCancelled('检查已停止，原始材料仍保留。')
                raise subprocess.TimeoutExpired(command, 180)
            try:
                known.update({p.pid:p for p in owner.children(recursive=True)})
                memory=0
                for child in known.values():
                    try:
                        if child.is_running(): memory+=child.memory_info().rss
                    except psutil.NoSuchProcess: pass
                if memory>ENGINE_RSS_LIMIT:
                    stop_tree()
                    raise DemoError('材料计算超过本机处理容量，检查已停止；请减少单次材料或联系管理员。')
            except psutil.NoSuchProcess:
                pass
            except psutil.AccessDenied:
                stop_tree()
                raise DemoError('无法核实检查进程的资源使用，已停止；请联系管理员。') from None
            try:
                stdout, stderr = process.communicate(timeout=.2)
                return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)
            except subprocess.TimeoutExpired:
                continue


def local_reviews(scan_id: str) -> review_store.ReviewStore:
    """Import legacy local results without pretending to authenticate platform users."""
    if not isinstance(scan_id, str) or not SAFE_SCAN_ID.fullmatch(scan_id):
        raise review_store.Unavailable("检查记录不存在。")
    job = job_store.get(RUNTIME_DIR, scan_id)
    if (job is not None and job.get('status') != 'succeeded') or not (RUNTIME_DIR / scan_id / 'report.html').is_file():
        raise review_store.Unavailable('本次检查尚未完成，暂不能保存复核或生成报告。')
    try:
        result = json.loads((RUNTIME_DIR / scan_id / 'result.json').read_text(encoding='utf-8'))
    except (OSError, ValueError):
        raise review_store.Unavailable("检查结果尚未保存，暂不能复核。") from None
    if result.get('scan_id') != scan_id:
        raise review_store.Conflict("检查结果编号不一致。")
    store = review_store.ReviewStore(RUNTIME_DIR)
    store.register(LOCAL_ACTOR, result)
    return store


def paperconan_executable() -> Path:
    configured = os.environ.get("PAPERCONAN_BIN")
    candidates = [
        Path(configured) if configured else None,
        Path(sys.executable).with_name("paperconan.exe" if os.name == "nt" else "paperconan"),
        ROOT / ".venv" / "Scripts" / "paperconan.exe",
        ROOT / ".venv" / "bin" / "paperconan",
    ]
    located = shutil.which("paperconan")
    if located:
        candidates.append(Path(located))
    for candidate in candidates:
        if candidate and candidate.is_file():
            return candidate.resolve()
    raise DemoError("检查服务未就绪，请联系管理员。")


def engine_version() -> str:
    try:
        completed = subprocess.run(
            [str(paperconan_executable()), "--version"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=15,
            check=True,
        )
        return completed.stdout.strip() or "paperconan"
    except (DemoError, OSError, subprocess.SubprocessError) as exc:
        return f"unavailable: {exc}"


def validate_upload(filename: str, payload: bytes, profile: str) -> str:
    if (not filename or filename != Path(filename).name
            or any(character in filename for character in '<>:"/\\|?*')
            or any(ord(character) < 32 for character in filename)
            or filename.endswith((".", " "))
            or re.fullmatch(r"(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])", Path(filename).stem)):
        raise DemoError("文件名不合法。")
    if Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise DemoError("不支持该文件类型。支持 CSV、TSV、Excel、PDF 和 DOCX。")
    if not payload:
        raise DemoError("文件内容为空。")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise DemoError(f"单个文件不能超过 {MAX_UPLOAD_BYTES // 1024 // 1024} MB。")
    if profile not in ALLOWED_PROFILES:
        raise DemoError("检查范围无效，请重新选择。")
    return filename


def validate_batch(uploads: list[tuple[str, bytes]], profile: str) -> None:
    if not uploads:
        raise DemoError("请先选择材料。")
    if len(uploads) > MAX_FILES:
        raise DemoError(f"一次最多选择 {MAX_FILES} 份材料。")
    seen: set[str] = set()
    total = 0
    for filename, payload in uploads:
        validate_upload(filename, payload, profile)
        if filename.casefold() in seen:
            raise DemoError(f"材料中有重名文件：{filename}。请修改文件名后重新选择。")
        seen.add(filename.casefold())
        total += len(payload)
    if total > MAX_BATCH_BYTES:
        raise DemoError("本次材料合计不能超过 100 MB。")


def _batch_label(filenames: list[str]) -> str:
    return filenames[0] if len(filenames) == 1 else f"{filenames[0]} 等 {len(filenames)} 份材料"


def _to_list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _finding_location(block: dict[str, Any], item: dict[str, Any]) -> str:
    file_name = str(item.get("file") or block.get("file") or "未知文件")
    sheet = str(item.get("sheet") or block.get("sheet") or "")
    rows = block.get("block", {}).get("rows") if isinstance(block.get("block"), dict) else None
    parts = [file_name]
    if sheet:
        parts.append(sheet)
    if rows:
        parts.append(f"行 {rows}")
    return " · ".join(parts)


def _display_number(value: Any) -> str:
    if isinstance(value, float):
        return f"{value:.8f}".rstrip("0").rstrip(".")
    return str(value)


def _friendly_benign(value: Any) -> str:
    message = str(value or "").replace("Methods", "研究方法").strip()
    if message and any("\u4e00" <= character <= "\u9fff" for character in message):
        return message
    return "请先确认是否属于单位换算、公式计算、共享对照、取整或同一数据的不同展示。"


def _focus_cells(
    kind: str,
    item: dict[str, Any],
    headers: list[Any],
    rows: list[Any],
    block: dict | None = None,
) -> list[dict[str, Any]]:
    if kind != "identical_after_rounding":
        return []
    displayed = item.get("rounded_to")
    if not isinstance(displayed, (int, float)) or not math.isfinite(displayed):
        return []
    # Engine example_values are rounded to four decimals, NOT source values.
    # Restrict reconstruction to its exact block; never search other tables by value.
    scope = (block or {}).get('block') or {}
    row_bounds = re.fullmatch(r'(\d+)-(\d+)', str(scope.get('rows', '')))
    col_bounds = re.fullmatch(r'(\d+)-(\d+)', str(scope.get('cols', '')))
    coordinates = {tuple(c) for c in item.get('example_cells', [])
                   if isinstance(c, (list, tuple)) and len(c) == 2 and all(isinstance(v, int) for v in c)}
    if not (row_bounds and col_bounds) and not coordinates:
        return []
    found: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        row_idx = row.get('row_idx')
        if not isinstance(row_idx, int):
            continue
        for index, value in enumerate(_to_list(row.get("values"))):
            if row_bounds and col_bounds:
                if not (int(row_bounds[1]) <= row_idx <= int(row_bounds[2])
                        and int(col_bounds[1]) <= index+1 <= int(col_bounds[2])):
                    continue
            elif (row_idx, index+1) not in coordinates:
                continue
            try:
                numeric = float(value)
            except (TypeError, ValueError):
                continue
            if math.isfinite(numeric) and 1e-9 < abs(numeric) < 100 and round(numeric, 1) == displayed:
                found.append(
                    {
                        "row_idx": row.get("row_idx", ""),
                        "column": str(headers[index]) if index < len(headers) else f"第 {index + 1} 列",
                        "original": numeric,
                        "displayed": displayed,
                    }
                )
                if len(found) >= 500:
                    return found
    return found


def _load_source_tables(source_path: Path, allowed_sheets: set[str] | None = None, cell_limit: int | None = None) -> dict[str, dict[str, list[Any]]]:
    """Bound source evidence independently of the engine preview; disclose truncation."""
    suffix = source_path.suffix.lower()
    tables: dict[str, dict[str, list[Any]]] = {}
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        try:
            with source_path.open("r", encoding="utf-8-sig", newline="") as handle:
                headers, rows, limited = statistical_checks._take_rows(iter(csv.reader(handle, delimiter=delimiter)), False, cell_limit)
        except (OSError, UnicodeError, csv.Error):
            return tables
        if headers:
            tables[source_path.stem] = {
                "headers": headers, "rows": rows, "limited": limited,
            }
            recover_headers(tables[source_path.stem])
        return tables

    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            from source_catalog import zip_guard

            zip_guard(source_path)
            workbook = load_workbook(source_path, data_only=True, read_only=True)
            remaining = statistical_checks.MAX_CELLS if cell_limit is None else cell_limit
            try:
                for worksheet in workbook.worksheets:
                    if allowed_sheets is not None and worksheet.title not in allowed_sheets:
                        continue
                    if (worksheet.max_column or 0) > statistical_checks.MAX_COLUMNS or remaining <= 0:
                        tables[worksheet.title] = {'headers': [], 'rows': [], 'limited': True}
                        continue
                    headers, rows, limited = statistical_checks._take_rows(iter(worksheet.iter_rows(values_only=True)), False, remaining)
                    remaining -= len(headers) + sum(len(r['values']) for r in rows)
                    if headers:
                        tables[worksheet.title] = {'headers': headers, 'rows': rows, 'limited': limited}
                        recover_headers(tables[worksheet.title])
            finally:
                workbook.close()
        except (ImportError, OSError, ValueError, BadZipFile, KeyError):
            return tables
    return tables


def _plain_explanation(
    kind: str, item: dict[str, Any], columns: list[str], n: int | float
) -> dict[str, Any]:
    left = columns[0] if columns else "第一组数据"
    right = columns[1] if len(columns) > 1 else "另一组数据"
    size = _display_number(n) if n else "多"
    common_next = ["确认两列是否由同一公式计算。", "核对单位、图注和原始记录。"]
    if kind == "constant_offset":
        offset = _display_number(item.get("offset", "固定数值"))
        return {
            "title": "固定差值",
            "summary": f"{right} 比 {left} 均高 {offset}，共 {size} 行。",
            "why": "两组本应独立变化的数据逐行保持完全相同的差值，规律过于整齐，因此需要确认是否为合法校正或派生列。",
            "next_steps": common_next,
        }
    if kind == "constant_ratio":
        ratio = _display_number(item.get("ratio", "固定倍数"))
        return {
            "title": "固定倍数",
            "summary": f"{right} 均为 {left} 的 {ratio} 倍，共 {size} 行。",
            "why": "整列保持完全一致的倍数关系，常见于单位换算或公式派生，也可能意味着两列并非独立测量。",
            "next_steps": common_next,
        }
    if kind == "exact_linear":
        slope = _display_number(item.get("slope", "k"))
        intercept = _display_number(item.get("intercept", "b"))
        return {
            "title": "固定换算",
            "summary": f"{right} 由 {left} 乘 {slope} 后加 {intercept} 得到，共 {size} 行。",
            "why": "真实独立测量通常会存在随机误差；完全贴合一条直线时，应先确认是否为明确记录的换算公式。",
            "next_steps": common_next,
        }
    if kind == "missing_last_digits":
        return {
            "title": "末位数字集中",
            "summary": f"共 {size} 个数值，部分末位数字未出现。",
            "why": "这可能来自仪器精度、人工取整或录入习惯，单独出现时证据较弱，需要结合测量精度判断。",
            "next_steps": [
                "确认仪器精度和保留的小数位。",
                "检查是否经过取整或格式调整。",
            ],
        }
    if kind == "rounded_to_half_or_int":
        return {'title': '记录精度集中',
                'summary': f'{left} 的 {size} 个记录中，{item.get("count_05", size)} 个数值的末位为 0 或 5。',
                'why': '仪器刻度或取整可以产生这种规律，需结合测量方法核对。',
                'next_steps': ['核对测量仪器精度。', '查看记录是否按固定刻度取整。']}
    if kind == "identical_after_rounding":
        rounded = _display_number(item.get("rounded_to", "同一值"))
        return {
            "title": "小数显示重复",
            "summary": f"{size} 个数值若保留 1 位小数均为 {rounded}，原始数值并不完全相同。",
            "why": "这通常只是显示精度造成的聚集，优先级较低；只有在论文把显示值当作不同独立结果解释时才值得继续核对。",
            "next_steps": [
                "查看原始文件中的完整小数位。",
                "确认表格使用了相同的显示精度。",
            ],
        }
    labels = {
        "equal_pairs": "相同数据",
        "progressions": "机械数列",
        "row_pairs": "重复行模式",
        "row_relations": "行间固定关系",
        "block_dups": "重复数据块",
        "cross_sheet": "跨表重复",
        "grim": "统计量不自洽",
    }
    title = labels.get(kind, labels.get(str(item.get("category")), "数值模式异常"))
    return {
        "title": title,
        "summary": f"共 {size} 个数值符合“{title}”。",
        "why": "该模式需要结合原始表头、实验条件和计算方法判断，当前只能视为待解释信号。",
        "next_steps": common_next,
    }


def _normalise_finding(
    block: dict[str, Any],
    item: dict[str, Any],
    category: str,
    index: int,
    source_tables: dict[str, dict[str, list[Any]]] | None = None,
) -> dict[str, Any]:
    evidence = item.get("evidence") if isinstance(item.get("evidence"), dict) else {}
    evidence_headers = _to_list(evidence.get("headers"))
    evidence_rows = _to_list(evidence.get("rows"))
    if evidence_rows and isinstance(evidence_rows[0], dict):
        first_values = _to_list(evidence_rows[0].get("values"))
        if [str(value) for value in first_values] == [str(value) for value in evidence_headers]:
            evidence_rows = evidence_rows[1:]
    kind = str(item.get("kind") or category)
    focus_headers = evidence_headers
    focus_rows = evidence_rows
    source_table = None
    sheet_name = str(item.get("sheet") or block.get("sheet") or "")
    if source_tables:
        source_table = source_tables.get(sheet_name)
        if source_table:
            focus_headers = _to_list(source_table.get("headers"))
            focus_rows = _to_list(source_table.get("rows"))
    focus_cells = _focus_cells(kind, item, focus_headers, focus_rows, block) if source_table else []
    # Engine highlights are absolute zero-based columns, while the UI uses
    # one-based positions within the displayed slice.
    col_offset = evidence.get('col_offset', 0)
    col_offset = col_offset if isinstance(col_offset, int) and col_offset >= 0 else 0
    if source_table and evidence_headers:
        end = col_offset + len(evidence_headers)
        if end <= len(focus_headers):
            evidence_headers = focus_headers[col_offset:end]
            by_row = {r['row_idx']: r for r in focus_rows}
            evidence_rows = [{**r, 'values': by_row[r['row_idx']]['values'][col_offset:end]}
                             for r in evidence_rows if r.get('row_idx') in by_row]
    columns = [item.get("col_a"), item.get("col_b")]
    columns = [str(value) for value in columns if value not in (None, "")]
    if not columns and item.get('col') not in (None,''):
        columns=[str(item['col'])]
        if source_table and isinstance(item.get('col_idx'),int) and 0<=item['col_idx']<len(focus_headers):
            columns=[str(focus_headers[item['col_idx']])]
    if not columns and source_table and isinstance(item.get('col_idx'), int):
        col_idx = item['col_idx']
        if 0 <= col_idx < len(focus_headers): columns = [str(focus_headers[col_idx])]
    if not columns and focus_cells:
        columns = list(dict.fromkeys(str(cell["column"]) for cell in focus_cells))
    rule = str(item.get("rule") or item.get("description") or kind)
    n = item.get("n") or item.get("n_cells") or item.get("overlap_n") or 0
    explanation = _plain_explanation(kind, item, columns, n)
    return {
        "id": f"FD-{index:04d}",
        "kind": kind,
        "category": category,
        "severity": str(item.get("severity") or item.get("raw_severity") or "info").lower(),
        "location": _finding_location(block, item),
        "file": str(item.get("file") or block.get("file") or ""),
        "sheet": str(item.get("sheet") or block.get("sheet") or ""),
        "columns": columns,
        "n": n,
        "rule": rule,
        "title_cn": explanation["title"],
        "summary_cn": explanation["summary"],
        "why_cn": explanation["why"],
        "next_steps": explanation["next_steps"],
        "profile_action": str(item.get("profile_action") or "kept"),
        "likely_benign": _friendly_benign(item.get("likely_benign")),
        "evidence": {
            "headers": evidence_headers,
            "highlight_cols": [c-col_offset+1 for c in _to_list(evidence.get("highlight_cols"))
                               if isinstance(c, int) and col_offset <= c < col_offset+len(evidence_headers)],
            "highlight_rows": _to_list(evidence.get("highlight_rows")),
            "rows": evidence_rows[:20],
            "focus_cells": focus_cells,
            "note_cn": (f'列出 {len(focus_cells)} / {n} 处原始数值。保留 1 位小数仅用于解释本条规律，不表示原文件采用此显示精度。'
                        if kind == 'identical_after_rounding' else ''),
            "truncated": evidence.get("truncated") if isinstance(evidence.get("truncated"), dict) else None,
        },
        "raw": {
            key: value
            for key, value in item.items()
            if key
            in {
                "offset",
                "ratio",
                "slope",
                "intercept",
                "rounded_to",
                "n_unique",
                "example_values",
            }
        },
    }


def normalise_scan(
    scan: dict[str, Any],
    scan_id: str,
    filename: str,
    profile: str,
    source_tables: dict[str, dict[str, list[Any]]] | None = None,
    file_tables: dict[str, dict[str, dict[str, list[Any]]]] | None = None,
) -> dict[str, Any]:
    finding_groups = (
        "relations",
        "equal_pairs",
        "progressions",
        "row_pairs",
        "row_relations",
        "within_col",
        "identical_after_rounding",
        "grim",
        "block_dups",
    )
    findings: list[dict[str, Any]] = []
    excluded=[]
    for block in _to_list(scan.get("relations_blocks")):
        if not isinstance(block, dict):
            continue
        for category in finding_groups:
            for item in _to_list(block.get(category)):
                if isinstance(item, dict):
                    identifier=(item.get('kind')=='arithmetic_progression' and re.fullmatch(
                        r'(?i)(?:subject(?:[ _]*(?:id|no\.?))?|participant[ _]*id|sample[ _]*id|patient[ _]*id|id|编号|序号|样本编号|受试者编号)',str(item.get('col','')).strip()))
                    if item.get('prefilter')=='drop' or identifier:
                        excluded.append({'file':str(item.get('file') or block.get('file') or ''),
                                         'sheet':str(item.get('sheet') or block.get('sheet') or ''),'kind':item.get('kind'),
                                         'columns':[str(item.get('col',''))],'n':item.get('n'),
                                         'reason_cn':'编号按顺序排列，不作为测量异常。' if identifier else '引擎已判为不应列入待核对的常见数据规律。'})
                        continue
                    findings.append(
                        _normalise_finding(
                            block,
                            item,
                            category,
                            len(findings) + 1,
                            (file_tables or {}).get(str(item.get("file") or block.get("file") or ""), source_tables),
                        )
                    )

    for item in _to_list(scan.get("cross_sheet_findings")):
        if isinstance(item, dict):
            findings.append(normalise_cross_finding(item, len(findings) + 1, file_tables or {}))

    severity_counts = Counter(item["severity"] for item in findings)
    coverage = scan.get("coverage") if isinstance(scan.get("coverage"), dict) else {}
    scan_stats = scan.get("scan_stats") if isinstance(scan.get("scan_stats"), dict) else {}
    sheets = _to_list(scan_stats.get("sheets"))
    numeric_cells = sum(int(item.get("numeric_cells") or 0) for item in sheets if isinstance(item, dict))
    return {
        "scan_id": scan_id,
        "source": filename,
        "profile": profile,
        "engine": str(scan.get("tool") or "paperconan"),
        "engine_version": str(scan.get("tool_version") or "unknown"),
        "schema_version": str(scan.get("schema_version") or "unknown"),
        "scan_status": str(scan.get("scan_status") or "unknown"),
        "summary": {
            "files": int(scan.get("n_files") or coverage.get("files_discovered") or 0),
            "sheets": int(coverage.get("sheets_succeeded") or len(sheets)),
            "blocks": int(coverage.get("blocks_analyzed") or 0),
            "numeric_cells": numeric_cells,
            "findings": len(findings),
            "severity": dict(severity_counts),
        },
        "coverage": coverage,
        "limitations": _to_list(coverage.get("limitations")),
        "errors": _to_list(scan.get("scan_errors")),
        "findings": findings,
        "excluded_signals": excluded,
        "artifacts": {
            "raw_scan": f"/artifact/{scan_id}/scan.json",
            "user_report": f"/artifact/{scan_id}/report.html",
            "result_json": f"/artifact/{scan_id}/result.json",
        },
        "notice": "筛查结果只提示需要人工核对的数据规律，不是对作者意图或科研不端的判断。",
    }


def _user_report_html(result: dict[str, Any]) -> str:
    def safe(value: Any) -> str:
        return html.escape(str(value), quote=True)

    summary = result["summary"]
    provenance_html = "".join(
        f"<p><strong>{safe(item['filename'])}</strong> · "
        + (f"公开材料：{safe(item['source_url'])} · 获取时间：{safe(item.get('retrieved_at', ''))}"
           if item.get("source_url") else "用户上传") + "</p>"
        for item in result.get("material_sources", [])
    )
    if provenance_html:
        provenance_html = f"<section class='item'><h2>材料来源</h2>{provenance_html}</section>"
    coverage_summary = result.get("coverage_summary") or {}
    material_rows = "".join(
        f"<tr><td>{safe(material['filename'])}</td><td>{safe({'read': '已读取', 'partial': '部分读取', 'failed': '未能读取', 'unread': '无可检查数据'}.get(material['status'], '待核对'))}</td>"
        f"<td>{safe(material.get('sheets', 0))}</td><td>{safe(material.get('numeric_cells', 0))}</td><td>{safe(material.get('reason_cn', ''))}</td></tr>"
        for material in result.get("materials", [])
    )
    coverage_html = ""
    if material_rows:
        coverage_html = (
            "<section class='item'><h2>材料读取情况</h2>"
            f"<p>{safe(coverage_summary.get('message_cn', ''))}</p>"
            "<div class='table-wrap'><table><thead><tr><th>材料</th><th>读取情况</th><th>数据表</th><th>数值</th><th>说明</th></tr></thead>"
            f"<tbody>{material_rows}</tbody></table></div>"
            + "".join(f"<p class='meta'>{safe(note)}</p>" for note in result.get("coverage_notes_cn", []))
            + "</section>"
        )
    statistical_html = ""
    if result.get("statistical_coverage"):
        statistical_rows = "".join(
            f"<tr><td>{safe(item['file'])} · {safe(item['sheet'])}</td><td>{safe(item['name'])}</td>"
            f"<td>{safe({'checked':'已核对','partial':'部分核对','insufficient':'未核对','unsupported':'暂不支持','failed':'未完成'}.get(item['status'],'未核对'))}</td>"
            f"<td>{item['checked_rows']} / {item['skipped_rows']}</td><td>{safe('；'.join(item['notes']))}</td></tr>"
            for item in result["statistical_coverage"]
        )
        assumptions = list(dict.fromkeys(item["assumption"] for item in result["statistical_coverage"]))
        statistical_html = ("<section class='item'><h2>统计核对范围</h2>"
                            + "".join(f"<p>用户确认的口径：{safe(text)}</p>" for text in assumptions)
                            + "<table><thead><tr><th>材料</th><th>核对内容</th><th>状态</th><th>已核对 / 跳过行数</th><th>说明</th></tr></thead>"
                            + f"<tbody>{statistical_rows}</tbody></table></section>")
    finding_sections: list[str] = []
    for finding in result["findings"]:
        evidence = finding.get("evidence") or {}
        headers = evidence.get("headers") or []
        rows = evidence.get("rows") or []
        focus_cells = evidence.get("focus_cells") or []
        table = ""
        comparison = evidence.get("comparison") or []
        if comparison:
            def cell_location(cell: dict[str, Any]) -> str:
                return f"{cell.get('file', '')} · {cell.get('sheet', '')} · 第 {cell.get('row_idx', '')} 行 · {cell.get('column', '')}"

            body = "".join(
                f"<tr><td>{safe(cell_location(pair['left']))}</td><td>{safe(pair['left']['value'])}</td>"
                f"<td>{safe(cell_location(pair['right']))}</td><td>{safe(pair['right']['value'])}</td></tr>"
                for pair in comparison
            )
            table = "<div class='table-wrap'><table><thead><tr><th>左侧位置</th><th>原始值</th><th>右侧位置</th><th>原始值</th></tr></thead>" f"<tbody>{body}</tbody></table></div>"
        elif focus_cells:
            body = "".join(
                "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                    safe(cell.get("row_idx", "")),
                    safe(cell.get("column", "")),
                    safe(cell.get("original", "")),
                    safe(cell.get("displayed", "")),
                )
                for cell in focus_cells
                if isinstance(cell, dict)
            )
            table = "<div class='table-wrap'><table><thead><tr><th>行号</th><th>列名</th><th>完整原始数值</th><th>保留1位小数后</th></tr></thead>" f"<tbody>{body}</tbody></table></div>"
        elif rows and headers:
            head = "".join(f"<th>{safe(column)}</th>" for column in headers)
            body = "".join(
                "<tr><td>{}</td>{}</tr>".format(
                    safe(row.get("row_idx", "")),
                    "".join(f"<td>{safe(value)}</td>" for value in row.get("values", [])),
                )
                for row in rows
                if isinstance(row, dict)
            )
            table = f"<div class='table-wrap'><table><thead><tr><th>行号</th>{head}</tr></thead><tbody>{body}</tbody></table></div>"
        checks = "".join(f"<li>{safe(item)}</li>" for item in finding.get("next_steps", [])[:2])
        table += f"<p class='meta'>{safe(evidence.get('note_cn', ''))}</p>" if evidence.get("note_cn") else ""
        finding_sections.append(
            "<section class='item'>"
            f"<div class='item-head'><h2>{safe(finding['title_cn'])}</h2><span>{safe({'high': '高', 'medium': '中', 'low': '低', 'info': '提示'}.get(finding['severity'], finding['severity']))}</span></div>"
            f"<p class='location'>{safe(finding['location'])}</p>"
            f"<p>{safe(finding['summary_cn'])}</p>"
            f"{table}"
            f"<h3>核对</h3><ul>{checks}</ul>"
            "</section>"
        )
    empty_text = "本次已检查范围内未发现需要核对的数据。" if result.get("scan_status") == "complete" else "当前已读取范围内没有可列出的结果；未读取或未完成检查的部分仍需核对。"
    findings_html = "".join(finding_sections) or f"<section class='item'><p>{empty_text}</p></section>"
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>科研数据筛查报告</title><style>
body{{margin:0;background:#f5f7fa;color:#1e293b;font:14px/1.65 "Microsoft YaHei",sans-serif}}main{{max-width:980px;margin:32px auto;padding:0 20px}}header,.item{{background:#fff;border:1px solid #dce4ee;border-radius:12px;padding:22px;margin-bottom:16px}}h1{{margin:0 0 6px;font-size:24px}}h2{{margin:0;font-size:18px}}h3{{margin:18px 0 6px;font-size:14px}}p{{margin:6px 0}}.meta{{color:#64748b}}.stats{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin-top:18px}}.stat{{background:#f4f7fb;border-radius:8px;padding:10px}}.stat b{{display:block;font-size:20px}}.item-head{{display:flex;justify-content:space-between;gap:12px}}.item-head span{{color:#36536f}}.location{{color:#64748b}}.table-wrap{{overflow:auto;margin-top:14px}}table{{width:100%;border-collapse:collapse;font-size:13px}}th,td{{padding:7px 9px;border:1px solid #dce4ee;text-align:right;white-space:nowrap}}th{{background:#f4f7fb}}th:first-child,td:first-child{{text-align:center}}ul{{margin:6px 0;padding-left:22px}}footer{{color:#64748b;padding:4px 2px 30px}}@media(max-width:700px){{.stats{{grid-template-columns:repeat(2,1fr)}}}}
</style></head><body><main><header><h1>科研数据筛查报告</h1><p class="meta">文件：{safe(result['source'])}</p><div class="stats"><div class="stat">文件<b>{safe(summary['files'])}</b></div><div class="stat">数据表<b>{safe(summary['sheets'])}</b></div><div class="stat">数值<b>{safe(summary['numeric_cells'])}</b></div><div class="stat">待核对<b>{safe(summary['findings'])}</b></div></div></header>{provenance_html}{coverage_html}{statistical_html}{findings_html}<footer>结果用于辅助核对，不作违规认定。</footer></main></body></html>"""


def _add_material_coverage(result: dict[str, Any], scan: dict[str, Any], filenames: list[str], pdf_info=None) -> None:
    stats = scan.get("scan_stats") or {}
    files = {str(item.get("file")): item for item in _to_list(stats.get("files")) if isinstance(item, dict)}
    sheets = [item for item in _to_list(stats.get("sheets")) if isinstance(item, dict)]
    coverage = scan.get("coverage") or {}
    limitations = [item for item in _to_list(coverage.get("limitations")) if isinstance(item, dict)]
    materials = []
    for filename in filenames:
        entry = files.get(filename, {})
        related = [sheet for sheet in sheets if sheet.get("file") == filename and "numeric_cells" in sheet]
        numeric = sum(int(sheet.get("numeric_cells") or 0) for sheet in related)
        file_errors = [item for item in _to_list(scan.get("scan_errors"))
                       if isinstance(item, dict) and item.get("file") == filename]
        affected = [item for item in limitations if item.get("file") == filename] + [item for item in file_errors if item.get("sheet")]
        failed = bool(entry.get("error")) or any(
            not item.get("sheet") for item in file_errors
        )
        if failed:
            status, reason = "failed", "文件未能读取，请确认可以正常打开；必要时将数据另存为 Excel 或 CSV。"
        elif not entry or not numeric:
            status, reason = "unread", "未提取到可检查的数值，请提供包含数据的表格。"
        elif affected:
            status, reason = "partial", "已读取部分数据，部分表格或检查未完成。"
        else:
            status, reason = "read", ""
        pdf=(pdf_info or {}).get(filename)
        if pdf and not numeric and not failed:
            if pdf['status']=='scanned':
                reason='PDF 没有可读取的文字（可能为扫描页或图片）；本版不做图片识别，请使用配套 Excel/CSV 数据。'
            elif pdf['status']=='text':
                reason=('已读取论文文字，但未取得独立数值单元格。图中数值或“均值（区间）”等复合文字未当作原始观察值扫描；请使用配套分列数据。')
            elif pdf['status']=='encrypted':reason='PDF 带密码保护或编辑限制，本版未读取文字；请使用未加密版本或配套表格。'
            elif pdf['status']=='limited':reason='论文文字超过本次读取范围，请拆分或使用配套表格。'
            else:reason='PDF 文字读取未完成；请使用可以直接读取的文字 PDF 或配套表格。'
        materials.append({"filename": filename, "status": status, "sheets": len(related),
                          "numeric_cells": numeric, "reason_cn": reason,**({'pdf_reading':pdf} if pdf else {})})
    engine_status = str(scan.get("scan_status") or "unknown")
    if not any(item["numeric_cells"] for item in materials) or engine_status == "failed":
        status, message = "failed", "未能取得可检查的数据，请查看材料读取情况后重新提交。"
    elif engine_status == "partial" or limitations or coverage.get("truncated") or any(item["status"] != "read" for item in materials):
        status, message = "partial", "本次只完成部分检查，未读取或未完成检查的部分仍需核对。"
    elif engine_status == "complete":
        status, message = "complete", "本次材料读取完成。"
    else:
        status, message = "unknown", "检查范围尚未确认，请查看材料读取情况。"
    result.update(materials=materials, coverage_summary={"status": status, "message_cn": message},
                  engine_scan_status=engine_status, scan_status=status)
    notes = []
    if int(coverage.get("blocks_skipped") or 0):
        notes.append(f"有 {coverage['blocks_skipped']} 个数据区域未完成检查。")
    if int(coverage.get("sheets_skipped") or 0):
        notes.append(f"有 {coverage['sheets_skipped']} 张数据表未完成检查。")
    if len(filenames) > 1 or len(sheets) > 1:
        notes.append("跨表同位置小数重复仅检查每张表前 200 行；整列重复按可读取的数值列检查。")
    result["coverage_notes_cn"] = notes


def scan_bytes(
    filename: str,
    payload: bytes,
    profile: str = "review",
    runtime_dir: Path | None = None,
) -> dict[str, Any]:
    return scan_files([(filename, payload)], profile, runtime_dir)


def scan_files(
    uploads: list[tuple[str, bytes]],
    profile: str = "review",
    runtime_dir: Path | None = None,
    scan_id: str | None = None,
    material_sources: list[dict[str, Any]] | None = None,
    extra_checks: list[str] | None = None,
    cancel_event: threading.Event | None = None,
) -> dict[str, Any]:
    validate_batch(uploads, profile)
    extra_checks = statistical_checks.selections(extra_checks)
    scan_id = scan_id or uuid.uuid4().hex[:12]
    if not SAFE_SCAN_ID.fullmatch(scan_id):
        raise DemoError("任务编号不合法。")
    runtime_root = (runtime_dir or RUNTIME_DIR).resolve()
    run_dir = runtime_root / scan_id
    input_dir = run_dir / "input"
    audit_dir = run_dir / "audit"
    input_dir.mkdir(parents=True, exist_ok=False)
    for filename, payload in uploads:
        (input_dir / filename).write_bytes(payload)

    command = [
        str(paperconan_executable()),
        "scan",
        str(input_dir),
        "--out",
        str(audit_dir),
        "--profile",
        profile,
    ]
    try:
        completed = run_engine(command, cancel_event)
    except subprocess.TimeoutExpired as exc:
        raise DemoError("处理时间过长，请缩小文件或改用表格文件后重试。") from exc
    scan_path = audit_dir / "scan.json"
    if completed.returncode != 0 and not scan_path.is_file():
        (run_dir / "internal.log").write_text(
            "\n".join(part for part in (completed.stdout, completed.stderr) if part),
            encoding="utf-8",
        )
        raise DemoError("本次筛查未完成。请确认文件可以正常打开且包含可读取的数据表，然后重试。")

    if not scan_path.is_file():
        raise DemoError("检查结果未生成，请重试。")
    with scan_path.open("r", encoding="utf-8") as handle:
        scan = json.load(handle)
    failed_files = {item.get("file") for item in _to_list(scan.get("scan_errors"))
                    if isinstance(item, dict) and not item.get("sheet")}
    successful_sheets: dict[str, set[str]] = {}
    for sheet in _to_list((scan.get("scan_stats") or {}).get("sheets")):
        if isinstance(sheet, dict) and "numeric_cells" in sheet:
            successful_sheets.setdefault(str(sheet.get("file")), set()).add(str(sheet.get("sheet")))
    file_tables={name: _load_source_tables(input_dir / name, successful_sheets.get(name, set()),
                                         statistical_checks.MAX_CELLS // len(uploads))
                 for name, _ in uploads if name not in failed_files}
    result = normalise_scan(
        scan,
        scan_id,
        _batch_label([name for name, _ in uploads]),
        profile,
        file_tables=file_tables,
    )
    result["sources"] = [name for name, _ in uploads]
    supplied_sources = {item["filename"]: item for item in material_sources or []}
    result["material_sources"] = [
        {**supplied_sources.get(name, {"kind": "upload"}), "filename": name,
         "sha256": hashlib.sha256(payload).hexdigest(), "size": len(payload)}
        for name, payload in uploads
    ]
    result["summary"]["files"] = len(uploads)
    pdf_info={};pdf_deadline=time.monotonic()+25
    for name,_ in uploads:
        if Path(name).suffix.lower()=='.pdf':
            if cancel_event is not None and cancel_event.is_set():raise ScanCancelled('检查已停止，原始材料仍保留。')
            if time.monotonic()>=pdf_deadline:
                pdf_info[name]={'status':'limited'};continue
            try:
                probe=subprocess.run([sys.executable,'-X','utf8',str(ROOT/'pdf_read_status.py'),str(input_dir/name)],
                                     capture_output=True,text=True,encoding='utf-8',timeout=max(.1,pdf_deadline-time.monotonic()),check=True)
                pdf_info[name]=json.loads(probe.stdout)
            except (subprocess.SubprocessError,ValueError):pdf_info[name]={'status':'unavailable'}
    _add_material_coverage(result, scan, result["sources"],pdf_info)
    result['evidence_limits'] = [{'file': name, 'sheet': sheet}
                                for name, tables in file_tables.items() for sheet, table in tables.items()
                                if table.get('limited')]
    if result['evidence_limits']:
        if result['scan_status'] != 'failed': result['scan_status'] = 'partial'
        result['coverage_notes_cn'].append('部分原值明细超过读取容量，未完整列出；请拆分材料后核对。')
    result["extra_checks"] = extra_checks
    result["statistical_coverage"] = []
    for filename, _ in uploads:
        if cancel_event is not None and cancel_event.is_set():
            raise ScanCancelled('检查已停止，原始材料仍保留。')
        additional, checked = statistical_checks.check_file(input_dir / filename, successful_sheets.get(filename, set()), extra_checks)
        for item in additional:
            item["id"] = f"FD-{len(result['findings']) + 1:04d}"
            result["findings"].append(item)
        result["statistical_coverage"].extend(checked)
    if any(item["status"] != "checked" for item in result["statistical_coverage"]):
        if result["scan_status"] != "failed":
            result["scan_status"] = "partial"
        result["coverage_notes_cn"].append("所选统计核对有未完成的内容，请查看“统计核对范围”。")
    result["summary"]["findings"] = len(result["findings"])
    result["summary"]["severity"] = dict(Counter(item["severity"] for item in result["findings"]))
    with (run_dir / "result.json").open("w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
    (run_dir / "report.html").write_text(_user_report_html(result), encoding="utf-8")
    return result


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _public_job(job: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in job.items()
        if key not in {"payload", "uploads", "runtime_dir", "material_sources"}
    }


def _update_job(job_id: str, **changes: Any) -> None:
    with JOB_LOCK:
        if job_id in JOBS:
            if JOB_CANCEL.get(job_id) is not None and JOB_CANCEL[job_id].is_set():
                if changes.get('status') == 'running':
                    changes['status'] = 'cancelling'
                elif changes.get('status') == 'succeeded':
                    changes.update(status='cancelled', result=None, error='检查已停止，原始材料仍保留。')
            JOBS[job_id].update(changes)
            job_store.save(JOBS[job_id]["runtime_dir"], JOBS[job_id])


def _run_scan_job(job_id: str) -> None:
    with JOB_LOCK:
        job = dict(JOBS[job_id])
    try:
        if JOB_CANCEL[job_id].is_set():
            raise ScanCancelled('检查已停止，原始材料仍保留。')
        _update_job(job_id, status="running", started_at=_utc_now())
        result = scan_files(
            job["uploads"],
            job["profile"],
            runtime_dir=job["runtime_dir"],
            scan_id=job_id,
            material_sources=job.get("material_sources"),
            extra_checks=job.get("extra_checks"),
            cancel_event=JOB_CANCEL[job_id],
        )
        if JOB_CANCEL[job_id].is_set():
            raise ScanCancelled('检查已停止，原始材料仍保留。')
        result.update(paper_id=job.get('paper_id',job_id),material_version=job.get('material_version',1),
                      parent_scan=job.get('parent_scan'))
        versioned=job['runtime_dir']/job_id/'versioned-result.pending'
        versioned.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        versioned.replace(job['runtime_dir']/job_id/'result.json')
        _update_job(
            job_id,
            status="succeeded",
            finished_at=_utc_now(),
            scan_id=result["scan_id"],
            scan_status=result["scan_status"],
            result=result,
        )
    except ScanCancelled as exc:
        _update_job(job_id, status='cancelled', finished_at=_utc_now(), error=str(exc), result=None)
    except Exception as exc:  # worker boundary records expected and unexpected failures
        message = str(exc) if isinstance(exc, DemoError) else "本次检查未完成，请确认材料可正常打开后重试。"
        print(f"Scan {job_id} failed: {exc}", file=sys.stderr)
        _update_job(job_id, status="failed", finished_at=_utc_now(), error=message)
    finally:
        with JOB_LOCK:
            JOBS[job_id].pop("uploads", None)
            JOB_CANCEL.pop(job_id, None)


def cancel_scan_job(job_id: str) -> dict:
    if not isinstance(job_id, str) or not SAFE_SCAN_ID.fullmatch(job_id):
        raise DemoError('检查编号无效。')
    with JOB_LOCK:
        job = JOBS.get(job_id)
        if job is None:
            saved = job_store.get(RUNTIME_DIR, job_id)
            if not saved: raise DemoError('检查记录不存在。')
            return _public_job(saved)
        if job['status'] in {'queued', 'running', 'cancelling'}:
            JOB_CANCEL[job_id].set()
            job.update(status='cancelling')
            job_store.save(job['runtime_dir'], job)
        return _public_job(dict(job))


def start_scan_job(
    filename: str,
    payload: bytes,
    profile: str = "review",
    runtime_dir: Path | None = None,
) -> dict[str, Any]:
    return start_scan_batch([(filename, payload)], profile, runtime_dir)


def start_scan_batch(
    uploads: list[tuple[str, bytes]],
    profile: str = "review",
    runtime_dir: Path | None = None,
    material_sources: list[dict[str, Any]] | None = None,
    extra_checks: list[str] | None = None,
    request_key: str | None = None,
    parent_scan: str | None = None,
) -> dict[str, Any]:
    root = (runtime_dir or RUNTIME_DIR).resolve()
    with SUBMISSION_LOCK:
        if request_key is not None and (not isinstance(request_key,str) or not re.fullmatch(r'[A-Za-z0-9_-]{8,100}',request_key)):
            raise DemoError('提交凭据无效。')
        signature = review_store.digest([[(name,hashlib.sha256(data).hexdigest()) for name,data in uploads],profile,extra_checks,material_sources,parent_scan])
        if request_key:
            prior=job_store.submission(root,request_key)
            if prior:
                if prior[0]!=signature: raise DemoError('同一次提交的材料已变化，请重新提交。')
                return get_scan_job(prior[1],root)
        with JOB_LOCK:
            active=sum(j.get('runtime_dir')==root and j.get('status') in {'queued','running','cancelling'} for j in JOBS.values())
        if active>=4: raise DemoError('已有多篇材料等待检查，请在当前任务结束后再提交。')
        lineage={}
        if parent_scan is not None:
            if not isinstance(parent_scan,str) or not SAFE_SCAN_ID.fullmatch(parent_scan): raise DemoError('原版本编号无效。')
            parent=job_store.get(root,parent_scan)
            if not parent or parent['status'] not in {'succeeded','failed','cancelled'}: raise DemoError('原版本尚未结束，不能创建新版本。')
            inherited=[]
            source_dir=root/parent_scan/'submitted'
            if not source_dir.is_dir(): source_dir=root/parent_scan/'input'
            for name,sha in parent.get('input_hashes',{}).items():
                if Path(name).name!=name: raise DemoError('原始材料位置无效。')
                try: payload=(source_dir/name).read_bytes()
                except OSError: raise DemoError('原始材料无法读取，请重新上传。') from None
                if hashlib.sha256(payload).hexdigest()!=sha: raise DemoError('原始材料校验未通过，不能建立新版本。')
                inherited.append((name,payload))
            if not inherited: raise DemoError('这条旧记录没有可核对的材料快照，请重新上传。')
            # Same-name replacements are explicit new-version uploads, never old-file writes.
            replacements={name for name,_ in uploads}
            merged=dict(inherited);merged.update(uploads);uploads=list(merged.items())
            new_sources=material_sources or []
            material_sources=[m for m in parent.get('material_sources',[]) if m.get('filename') not in replacements]+new_sources
            paper=parent.get('paper_id',parent_scan)
            lineage={'paper_id':paper,'parent_scan':parent_scan,'material_version':max(j.get('material_version',1) for j in job_store.family(root,paper))+1}
        return _start_scan_batch(uploads,profile,root,material_sources,extra_checks,
                                 (request_key,signature) if request_key else None,lineage)


def _start_scan_batch(uploads,profile,runtime_dir,material_sources,extra_checks,submission,lineage):
    validate_batch(uploads, profile)
    extra_checks = statistical_checks.selections(extra_checks)
    filenames = [name for name, _ in uploads]
    job_id = uuid.uuid4().hex[:12]
    job = {
        "job_id": job_id,
        "status": "queued",
        "filename": _batch_label(filenames),
        "filenames": filenames,
        "file_count": len(filenames),
        "profile": profile,
        "created_at": _utc_now(),
        "started_at": None,
        "finished_at": None,
        "scan_id": job_id,
        "result": None,
        "error": None,
        "uploads": uploads,
        "material_sources": material_sources or [],
        "extra_checks": extra_checks,
        "runtime_dir": (runtime_dir or RUNTIME_DIR).resolve(),
        "paper_id": job_id,
        "material_version": 1,
        **lineage,
    }
    submitted = job['runtime_dir'] / job_id / 'submitted'
    submitted.mkdir(parents=True, exist_ok=False)
    for filename, payload in uploads:
        (submitted / filename).write_bytes(payload)
    job['input_hashes'] = {name: hashlib.sha256(payload).hexdigest() for name, payload in uploads}
    with JOB_LOCK:
        job_store.save(job["runtime_dir"], job,submission)
        JOBS[job_id] = job
        JOB_CANCEL[job_id] = threading.Event()
    queued = _public_job(dict(job))
    JOB_EXECUTOR.submit(_run_scan_job, job_id)
    return queued


def get_scan_job(job_id: str, runtime_dir: Path | None = None) -> dict[str, Any]:
    if not SAFE_SCAN_ID.fullmatch(job_id):
        raise DemoError("任务编号不合法。")
    with JOB_LOCK:
        job = JOBS.get(job_id)
        if job and (runtime_dir is None or job["runtime_dir"] == runtime_dir.resolve()):
            return _public_job(dict(job))
    root = (runtime_dir or RUNTIME_DIR).resolve()
    saved = job_store.get(root, job_id)
    if not saved:
        raise DemoError("筛查记录不存在。")
    saved["result"] = None
    if saved["status"] == "succeeded" and SAFE_SCAN_ID.fullmatch(saved.get("scan_id") or ""):
        try:
            saved["result"] = json.loads((root / saved["scan_id"] / "result.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            saved.update(status="failed", error="这条记录的结果文件无法读取，请重新提交材料。")
    return saved


def list_scan_jobs(limit: int = 50, runtime_dir: Path | None = None) -> list[dict[str, Any]]:
    return job_store.list_jobs((runtime_dir or RUNTIME_DIR).resolve(), limit)


def restore_scan_jobs(runtime_dir: Path | None = None) -> None:
    root = (runtime_dir or RUNTIME_DIR).resolve()
    job_store.recover_interrupted(root, _utc_now())
    known = job_store.known_scan_ids(root)
    # Old releases saved reports but not a task index. Import only valid completed artifacts.
    for result_path in root.glob("*/result.json"):
        scan_id = result_path.parent.name
        if (not SAFE_SCAN_ID.fullmatch(scan_id) or scan_id in known
                or not (result_path.parent / "report.html").is_file()):
            continue
        try:
            result = json.loads(result_path.read_text(encoding="utf-8"))
            if result.get("scan_id") != scan_id or not isinstance(result.get("summary"), dict):
                continue
            filenames = result.get("sources") or [result["source"]]
            created = datetime.fromtimestamp(result_path.stat().st_mtime, timezone.utc).isoformat()
            job_store.save(root, {"job_id": scan_id, "filename": result["source"],
                                 "filenames": filenames, "file_count": len(filenames),
                                 "profile": result.get("profile", "review"), "status": "succeeded",
                                 "scan_status": result.get("scan_status") or (result.get("coverage_summary") or {}).get("status", "unknown"),
                                 "created_at": created, "started_at": created, "finished_at": created,
                                 "scan_id": scan_id, "error": None})
        except (OSError, ValueError, KeyError, TypeError) as exc:
            print(f"Could not import scan {scan_id}: {exc}", file=sys.stderr)


def decode_upload_body(body: dict[str, Any]) -> tuple[str, bytes, str]:
    filename = str(body.get("filename") or "")
    profile = str(body.get("profile") or "review")
    encoded = str(body.get("content_base64") or "")
    if encoded.startswith("data:") and "," in encoded:
        encoded = encoded.split(",", 1)[1]
    try:
        payload = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise DemoError("上传内容不是有效的 Base64 数据。") from exc
    return filename, payload, profile


def decode_batch_body(body: dict[str, Any]) -> tuple[list[tuple[str, bytes]], str]:
    profile = str(body.get("profile") or "review")
    if "files" not in body:
        filename, payload, profile = decode_upload_body(body)
        uploads = [(filename, payload)]
    else:
        files = body["files"]
        if not isinstance(files, list) or not 1 <= len(files) <= MAX_FILES:
            raise DemoError("请选择 1 至 20 份材料。")
        uploads = []
        for item in files:
            if not isinstance(item, dict):
                raise DemoError("材料格式不正确，请重新选择。")
            filename, payload, _ = decode_upload_body({**item, "profile": profile})
            uploads.append((filename, payload))
    validate_batch(uploads, profile)
    return uploads, profile


def material_store() -> LocalMaterialStore:
    return LocalMaterialStore(RUNTIME_DIR.parent / (RUNTIME_DIR.name + "-materials"))


def decode_submission(body: dict[str, Any]) -> tuple[list[tuple[str, bytes]], str, list[dict[str, Any]]]:
    """Resolve server-owned snapshots; ignore browser-provided source descriptions."""
    if not isinstance(body.get("files"), list) or not any(isinstance(item, dict) and "source_id" in item for item in body["files"]):
        uploads, profile = decode_batch_body(body)
        return uploads, profile, []
    if not 1 <= len(body["files"]) <= MAX_FILES:
        raise DemoError("请选择 1 至 20 份材料。")
    uploads, sources = [], []
    profile = str(body.get("profile") or "review")
    for item in body["files"]:
        if not isinstance(item, dict):
            raise DemoError("材料格式不正确，请重新选择。")
        if "source_id" in item:
            if set(item) != {"source_id"}:
                raise DemoError("已获取材料不能同时提交修改内容，请重新选择。")
            metadata, payload = material_store().load(LOCAL_ACTOR, item["source_id"])
            filename = metadata["filename"]
            sources.append({**metadata["provenance"], "filename": filename,
                            "source_id": metadata["source_id"], "sha256": metadata["sha256"]})
        else:
            filename, payload, _ = decode_upload_body({**item, "profile": profile})
        uploads.append((filename, payload))
    validate_batch(uploads, profile)
    return uploads, profile, sources


class DemoHandler(BaseHTTPRequestHandler):
    server_version = "PaperAuditMVP/0.1"

    def log_message(self, fmt: str, *args: Any) -> None:
        sys.stdout.write(f"[http] {self.address_string()} - {fmt % args}\n")

    def _json(self, status: int, data: dict[str, Any]) -> None:
        payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(payload)

    def _file(self, path: Path, content_type: str | None = None) -> None:
        if not path.is_file():
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        payload = path.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(payload)

    def _request_json(self) -> dict[str, Any]:
        if self.headers.get("Content-Type", "").split(";", 1)[0].strip().lower() != "application/json":
            raise DemoError("请从本机工作台提交材料。")
        length = int(self.headers.get("Content-Length") or "0")
        if length <= 0 or length > (MAX_BATCH_BYTES * 4 // 3 + 1024 * 1024):
            raise DemoError("请求体为空或过大。")
        body = json.loads(self.rfile.read(length).decode("utf-8"))
        if not isinstance(body, dict):
            raise DemoError("提交的数据格式不正确。")
        return body

    def _local_request(self) -> bool:
        """Block DNS rebinding and cross-site invocation of the unauthenticated local adapter."""
        port = self.server.server_address[1]
        allowed = {f"127.0.0.1:{port}", f"localhost:{port}", f"[::1]:{port}"}
        host = self.headers.get("Host", "").lower()
        origin = self.headers.get("Origin")
        if host not in allowed or (origin is not None and origin != f"http://{host}"):
            self._json(HTTPStatus.FORBIDDEN, {"ok": False, "error": "请从本机工作台访问。"})
            return False
        return True

    def do_GET(self) -> None:  # noqa: N802
        if not self._local_request():
            return
        path = unquote(urlparse(self.path).path)
        if path in {"/", "/index.html"}:
            self._file(STATIC_DIR / "index.html", "text/html; charset=utf-8")
            return
        if path == "/public-materials.js":
            self._file(STATIC_DIR / "public-materials.js", "text/javascript; charset=utf-8")
            return
        if path == "/review-ui.js":
            self._file(STATIC_DIR / "review-ui.js", "text/javascript; charset=utf-8")
            return
        if path == '/reference-ui.js':
            self._file(STATIC_DIR/'reference-ui.js','text/javascript; charset=utf-8')
            return
        if path == '/coverage-labels.js':
            self._file(STATIC_DIR/'coverage-labels.js','text/javascript; charset=utf-8')
            return
        if path.startswith('/api/reference/'):
            scan=path.removeprefix('/api/reference/')
            try:
                local_reviews(scan)
                self._json(HTTPStatus.OK,{'ok':True,'comparisons':reference_comparison.ReferenceComparisonService(RUNTIME_DIR).reports(LOCAL_ACTOR,scan)})
            except review_store.ReviewError as exc:self._json(exc.status,{'ok':False,'error':str(exc)})
            return
        if path == '/content-ui.js':
            self._file(STATIC_DIR / 'content-ui.js', 'text/javascript; charset=utf-8')
            return
        if path == '/versions-ui.js':
            self._file(STATIC_DIR / 'versions-ui.js', 'text/javascript; charset=utf-8')
            return
        if path == '/finding-ai.js':
            self._file(STATIC_DIR / 'finding-ai.js', 'text/javascript; charset=utf-8')
            return
        if path.startswith('/api/explanations/'):
            parts=path.strip('/').split('/')
            if len(parts)!=4 or not SAFE_SCAN_ID.fullmatch(parts[2]) or not re.fullmatch(r'FD-[0-9]{4}',parts[3]):
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            try:
                local_reviews(parts[2])
                answer=finding_explanations.FindingExplanationService(RUNTIME_DIR).view(LOCAL_ACTOR,parts[2],parts[3])
                self._json(HTTPStatus.OK,{'ok':True,'ai':answer})
            except review_store.ReviewError as exc:
                self._json(exc.status,{'ok':False,'error':str(exc)})
            return
        if path.startswith('/api/versions/'):
            try:
                job = get_scan_job(path.removeprefix('/api/versions/'))
                versions=job_store.family(RUNTIME_DIR,job.get('paper_id',job['job_id']))
                self._json(HTTPStatus.OK,{'ok':True,'versions':[{k:j.get(k) for k in ['job_id','parent_scan','material_version','status','created_at']} for j in versions]})
            except DemoError as exc:
                self._json(HTTPStatus.NOT_FOUND,{'ok':False,'error':str(exc)})
            return
        if path.startswith('/api/content/'):
            scan = path.removeprefix('/api/content/')
            try:
                local_reviews(scan)
                self._json(HTTPStatus.OK, {'ok':True, 'content':agent_service.AgentService(RUNTIME_DIR).view(LOCAL_ACTOR,scan)})
            except review_store.ReviewError as exc:
                self._json(exc.status, {'ok':False,'error':str(exc)})
            return
        if path == "/api/checks":
            self._json(HTTPStatus.OK, {"ok": True, "checks": [
                {"key": key, "name": rule['name'], "assumption": rule['assumption']}
                for key, rule in statistical_checks.RULES.items()]})
            return
        if path.startswith('/api/reviews/') or path.startswith('/api/reports/'):
            parts = path.strip('/').split('/')
            try:
                if len(parts) == 3 and parts[1] == 'reviews':
                    store = local_reviews(parts[2])
                    self._json(HTTPStatus.OK, {'ok': True, 'review': store.view(LOCAL_ACTOR, parts[2]),
                                              'reports': store.reports(LOCAL_ACTOR, parts[2])})
                elif len(parts) == 4 and parts[1] == 'reports':
                    store = review_store.ReviewStore(RUNTIME_DIR)
                    self._file(report_exports.download_path(store, LOCAL_ACTOR, parts[2], parts[3]))
                else:
                    self.send_error(HTTPStatus.NOT_FOUND)
            except review_store.ReviewError as exc:
                self._json(exc.status, {'ok': False, 'error': str(exc)})
            return
        if path == "/api/health":
            version = engine_version()
            self._json(
                HTTPStatus.OK,
                {
                    "ok": not version.startswith("unavailable:"),
                    "service": "PaperAudit MVP",
                    "runtime_id":hashlib.sha256(RUNTIME_DIR.resolve().as_posix().casefold().encode('utf-8')).hexdigest()[:16],
                    "engine": version,
                    "max_upload_mb": MAX_UPLOAD_BYTES // 1024 // 1024,
                    "max_batch_mb": MAX_BATCH_BYTES // 1024 // 1024,
                    "max_files": MAX_FILES,
                    "formats": sorted(ALLOWED_EXTENSIONS),
                },
            )
            return
        if path == "/api/jobs":
            self._json(HTTPStatus.OK, {"ok": True, "jobs": list_scan_jobs()})
            return
        if path.startswith("/api/jobs/"):
            parts = path.strip("/").split("/")
            if len(parts) != 3:
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            try:
                self._json(HTTPStatus.OK, {"ok": True, "job": get_scan_job(parts[2])})
            except DemoError as exc:
                self._json(HTTPStatus.NOT_FOUND, {"ok": False, "error": str(exc)})
            return
        if path == "/samples/fixed_shift_demo.csv":
            self._file(SAMPLES_DIR / "fixed_shift_demo.csv", "text/csv; charset=utf-8")
            return
        if path.startswith("/artifact/"):
            parts = path.strip("/").split("/")
            if len(parts) != 3 or not SAFE_SCAN_ID.fullmatch(parts[1]) or parts[2] not in {"scan.json", "report.html", "result.json"}:
                self.send_error(HTTPStatus.NOT_FOUND)
                return
            if parts[2] in {"report.html", "result.json"}:
                artifact_path = RUNTIME_DIR / parts[1] / parts[2]
            else:
                artifact_path = RUNTIME_DIR / parts[1] / "audit" / parts[2]
            self._file(artifact_path)
            return
        self.send_error(HTTPStatus.NOT_FOUND)

    def do_POST(self) -> None:  # noqa: N802
        if not self._local_request():
            return
        path = urlparse(self.path).path
        if path not in {"/api/scan", "/api/jobs", "/api/materials/discover", "/api/materials/import", "/api/reviews", "/api/reports", "/api/jobs/cancel", '/api/jobs/retry', '/api/content/prepare', '/api/content/start', '/api/content/quote', '/api/content/cancel', '/api/explanations/quote', '/api/explanations/run','/api/reference/preview','/api/reference/compare'}:
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        try:
            body = self._request_json()
            if path.startswith('/api/reference/'):
                scan=body.get('scan_id');local_reviews(scan)
                service=reference_comparison.ReferenceComparisonService(RUNTIME_DIR)
                if path.endswith('/preview'):
                    filename,payload,_=decode_upload_body(body);validate_upload(filename,payload,'review')
                    answer=service.upload(LOCAL_ACTOR,scan,body.get('source_file'),filename,payload)
                else:answer=service.compare(LOCAL_ACTOR,scan,body.get('reference_id'),body.get('source_sheet'),body.get('reference_sheet'),
                            body.get('keys'),body.get('columns'),body.get('confirm_order'))
                self._json(HTTPStatus.OK,{'ok':True,'reference':answer});return
            if path.startswith('/api/explanations/'):
                scan=body.get('scan_id'); finding=body.get('finding_id')
                local_reviews(scan)
                service=finding_explanations.FindingExplanationService(RUNTIME_DIR)
                if path.endswith('/quote'):
                    answer=service.quote(LOCAL_ACTOR,scan,finding)
                else:
                    answer=service.run(LOCAL_ACTOR,scan,finding,body.get('quote_id'),
                                       body.get('request_key'),body.get('confirmed'))
                self._json(HTTPStatus.OK,{'ok':True,'ai':answer})
                return
            if path.startswith('/api/content/'):
                scan = body.get('scan_id'); local_reviews(scan)
                service = agent_service.AgentService(RUNTIME_DIR)
                if path.endswith('/prepare'): answer = service.prepare(LOCAL_ACTOR,scan,body.get('roles'))
                elif path.endswith('/start'): answer = service.start(LOCAL_ACTOR,scan,body.get('quote_id'),body.get('request_key'),body.get('confirmed'))
                elif path.endswith('/quote'): answer = service.quote(LOCAL_ACTOR,scan)
                else: answer = service.cancel(LOCAL_ACTOR,scan)
                self._json(HTTPStatus.OK,{'ok':True,'content':answer})
                return
            if path == '/api/jobs/cancel':
                self._json(HTTPStatus.OK, {'ok': True, 'job': cancel_scan_job(body.get('job_id'))})
                return
            if path == '/api/jobs/retry':
                parent=get_scan_job(body.get('parent_scan') or '')
                if parent['status'] not in {'failed','cancelled'}: raise DemoError('只有中断或未完成的检查可以重试。')
                job=start_scan_batch([],parent['profile'],extra_checks=parent.get('extra_checks',[]),request_key=body.get('request_key'),parent_scan=parent['job_id'])
                self._json(HTTPStatus.ACCEPTED,{'ok':True,'job':job})
                return
            if path in {'/api/reviews', '/api/reports'}:
                scan_id = body.get('scan_id')
                store = local_reviews(scan_id)
                if path == '/api/reviews':
                    view = store.save(LOCAL_ACTOR, scan_id, body.get('finding_id'), body.get('state'),
                                      body.get('note'), body.get('revision'), body.get('request_key'))
                    self._json(HTTPStatus.OK, {'ok': True, 'review': view})
                else:
                    with agent_service.LOCK, finding_explanations.LOCK, reference_comparison.LOCK:
                        content = agent_service.AgentService(RUNTIME_DIR).report_context(LOCAL_ACTOR,scan_id)
                        ai = finding_explanations.FindingExplanationService(RUNTIME_DIR).report_context(LOCAL_ACTOR,scan_id)
                        reference=reference_comparison.ReferenceComparisonService(RUNTIME_DIR).report_context(LOCAL_ACTOR,scan_id)
                        snap = store.freeze(LOCAL_ACTOR, scan_id, body.get('revision'), body.get('request_key'),content,finding_ai=ai,reference_review=reference)
                    exported = report_exports.export_snapshot(store, LOCAL_ACTOR, snap)
                    self._json(HTTPStatus.OK, {'ok': True, 'report': exported})
                return
            if path == "/api/materials/discover":
                self._json(HTTPStatus.OK, {"ok": True, "discovery": public_materials.discover(body.get("reference", ""))})
                return
            if path == "/api/materials/import":
                if body.get("confirmed") is not True:
                    raise DemoError("请先确认材料属于本次论文。")
                source = public_materials.import_material(body.get("url", ""), material_store(), LOCAL_ACTOR)
                self._json(HTTPStatus.OK, {"ok": True, "material": source})
                return
            uploads, profile, sources = decode_submission(body)
            extra_checks = statistical_checks.selections(body.get("extra_checks"))
            if path == "/api/jobs":
                job = start_scan_batch(uploads, profile, material_sources=sources, extra_checks=extra_checks,
                                       request_key=body.get('request_key'),parent_scan=body.get('parent_scan'))
                self._json(HTTPStatus.ACCEPTED, {"ok": True, "job": job})
            else:
                result = scan_files(uploads, profile, material_sources=sources, extra_checks=extra_checks)
                self._json(HTTPStatus.OK, {"ok": True, "result": result})
        except review_store.ReviewError as exc:
            self._json(exc.status, {'ok': False, 'error': str(exc)})
        except (DemoError, MaterialError, statistical_checks.CannotCheck, BudgetError, ModelError) as exc:
            self._json(HTTPStatus.BAD_REQUEST, {"ok": False, "error": str(exc)})
        except (json.JSONDecodeError, UnicodeDecodeError):
            self._json(HTTPStatus.BAD_REQUEST, {"ok": False, "error": "提交的内容无法读取。"})
        except Exception:  # pragma: no cover - defensive server boundary
            print("Request failed; details suppressed to protect material.", file=sys.stderr)
            self._json(HTTPStatus.INTERNAL_SERVER_ERROR, {"ok": False, "error": "本次操作未完成，请稍后重试。"})


def run_server(host: str = "127.0.0.1", port: int = 8765) -> None:
    if os.environ.get("PAPER_AUDIT_PLATFORM_MODE", "local") != "local":
        raise DemoError("佰茵云对接尚未配置，不能以本地身份代替平台登录。")
    if host != "localhost" and not ipaddress.ip_address(host).is_loopback:
        raise DemoError("本地版本仅允许本机访问；平台身份接入前不能对外开放。")
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    with RuntimeLease(RUNTIME_DIR):
        server = ThreadingHTTPServer((host, port), DemoHandler)
        try:
            restore_scan_jobs()
            agent_service.AgentService(RUNTIME_DIR).recover(LOCAL_ACTOR)
            print(f"PaperAudit: http://{host}:{port}")
            print(f"Engine: {engine_version()}")
            print("按 Ctrl+C 停止服务。")
            server.serve_forever()
        except KeyboardInterrupt:
            print("\n服务已停止。")
        finally:
            server.server_close()
            # Keep exclusive ownership until in-flight local work has stopped.
            for event in JOB_CANCEL.values(): event.set()
            agent_service.AgentService(RUNTIME_DIR).stop_running(LOCAL_ACTOR)
            JOB_EXECUTOR.shutdown(wait=True)
            agent_service.EXECUTOR.shutdown(wait=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="论文科研数据异常筛查 MVP")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args()
    run_server(args.host, args.port)
